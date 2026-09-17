"""
A tabela consolidada em XLSX, para juntada aos autos.

§16: o arquivo exportado reproduz EXATAMENTE os totais aprovados. Não é
um resumo nem uma versão simplificada — é a mesma consolidação, com as
mesmas parcelas, num formato que o processo administrativo aceita.

Três abas, com o mesmo desenho da planilha que a prefeitura já produzia
à mão — porque quem vai conferir já sabe ler aquela:

- **Consolidado**: uma coluna por secretaria e o total geral;
- **Ocorrências**: cada parcela com o arquivo e o DFD de onde veio;
- **Resumo**: o que entrou, o que não foi lido, o que ficou em conflito.

A aba de ocorrências é a que sustenta o §12. Sem ela o total seria uma
afirmação; com ela, alguém refaz a conta linha a linha.
"""

from __future__ import annotations

import io
from decimal import Decimal

from .consolidacao import Consolidado, conferir


def _numero(valor: Decimal) -> float | int:
    """
    Célula numérica de verdade, não texto.

    Número como texto no XLSX não soma, não ordena e não fecha conta —
    e a planilha existe justamente para alguém conferir a soma.
    """
    inteiro = valor.to_integral_value()
    return int(inteiro) if valor == inteiro else float(valor)


def planilha_consolidada(consolidado: Consolidado) -> bytes:
    """O XLSX completo, pronto para anexar ao processo."""
    import openpyxl

    wb = openpyxl.Workbook()
    secretarias = list(consolidado.secretarias)

    aba = wb.active
    aba.title = "Consolidado"
    aba.append(["Código do Item", "Descrição e Especificação", "Unidade"]
               + secretarias + ["Total Geral"])
    for linha in consolidado.linhas:
        aba.append(
            [linha.codigo, linha.descricao, linha.unidade]
            + [_numero(linha.por_secretaria[s]) if s in linha.por_secretaria else None
               for s in secretarias]
            + [_numero(linha.total)]
        )

    ocorrencias = wb.create_sheet("Ocorrências")
    ocorrencias.append(["Secretaria", "Código do Item", "Descrição", "Unidade",
                        "Quantidade", "DFD", "Arquivo de Origem"])
    for linha in consolidado.linhas:
        for origem in linha.origens:
            ocorrencias.append([
                origem.secretaria, linha.codigo,
                origem.descricao or linha.descricao, linha.unidade,
                _numero(origem.quantidade), origem.dfd, origem.arquivo,
            ])

    laudo = conferir(consolidado)
    resumo = wb.create_sheet("Resumo")
    resumo.append(["Indicador", "Valor"])
    resumo.append(["Itens consolidados", len(consolidado.linhas)])
    resumo.append(["Secretarias somadas", len(secretarias)])
    resumo.append(["Ocorrências rastreadas", laudo.total_de_ocorrencias])
    resumo.append(["A soma fecha", "sim" if laudo.fecha else "NÃO"])
    for arquivo in consolidado.nao_lidos:
        resumo.append(["Arquivo não lido", arquivo])
    for arquivo in consolidado.duplicatas:
        resumo.append(["Arquivo duplicado, ignorado", arquivo])
    for conflito in consolidado.conflitos:
        resumo.append([f"Conflito de unidade no item {conflito.codigo}",
                       ", ".join(conflito.unidades)])

    memoria = io.BytesIO()
    wb.save(memoria)
    return memoria.getvalue()


__all__ = ["planilha_consolidada"]
