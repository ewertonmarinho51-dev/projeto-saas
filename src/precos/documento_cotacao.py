"""
Documentos da cotação: Relatório de Cotação (compacto) e Relatório
Analítico Integral, em DOCX — e em PDF pelo conversor que o projeto já
usa (`export._docx_em_pdf`, LibreOffice).

POR QUE UM RENDERER PRÓPRIO, E NÃO MAIS MARKDOWN → `export.gerar_pdf`

O PDF de 907 páginas mostrou o limite do caminho genérico: tabelas de 9 a
11 colunas em A4 retrato, larguras proporcionais ao texto, palavras
partidas no cabeçalho ("Méto/do", "Situaç/ão", "Unitá/rio") e todas as
referências dentro do documento principal. Aqui cada tabela tem colunas
PROJETADAS para o conteúdo (largura fixa, cabeçalho que se repete,
linha que não se parte), e o volume vai para o anexo, não para o
relatório que se lê.

O que NÃO mudou: o conteúdo vem de `relatorio_view.CotacaoReportView`,
que lê o que o motor gravou. Nada aqui calcula preço; a única conta é a
formatação de Decimal em texto.

Tabelas e parágrafos são escritos direto em WordprocessingML: é o que
mantém o anexo de milhares de linhas em segundos (o `Table.cell` do
python-docx é quadrático) e dá controle exato de largura, sombreamento,
quebra e hyperlink.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from decimal import Decimal
from xml.sax.saxutils import escape

from .estatistica import brl, pct
from .relatorio_view import (NAO_INFORMADO, CotacaoReportView, ItemView,
                             ReferenciaView)

# Versão do LAYOUT. Entra no nome do registro de emissão e na chave de
# cache da tela: mudar o desenho muda o arquivo.
VERSAO_RENDERER = "cotacao-2026-10-17.revisao02"

# --- tokens visuais GovConnect (sem cores nem marcas de terceiros) ------
FONTE = "DejaVu Sans"
FONTE_TABELA = "DejaVu Sans Condensed"
AZUL = "1B4F8A"          # paleta institucional (.streamlit/config.toml)
TINTA = "172033"
SECUNDARIO = "4B5565"
BORDA = "D5DDE8"
FAIXA = "EEF3F9"
CABECALHO_TABELA = "E4ECF6"
ALERTA_FUNDO = "FFF4E5"
ALERTA_TINTA = "7A4A00"
EXCLUIDA = "9F2F2D"
BRANCO = "FFFFFF"

CM = 567                 # twips por centímetro
CORPO_PT = 9.5
TABELA_PT = 8.5
NOTA_PT = 8


class ErroDeRenderizacao(RuntimeError):
    """O documento não pôde ser convertido — mensagem para a tela."""


@dataclass(frozen=True)
class Coluna:
    titulo: str
    largura_cm: float
    alinhar: str = "left"      # left | right | center


# ---------------------------------------------------------------------------
# Formatação (texto a partir do que a projeção trouxe)
# ---------------------------------------------------------------------------
def moeda(valor: Decimal | None) -> str:
    return "—" if valor is None else brl(valor)


def quantidade(valor: Decimal | None) -> str:
    """1500 → "1.500"; 2.5 → "2,5"; sem notação científica."""
    if valor is None:
        return "—"
    texto = format(valor.normalize(), "f")
    inteiro, _, fracao = texto.partition(".")
    sinal = "-" if inteiro.startswith("-") else ""
    inteiro = inteiro.lstrip("-")
    grupos = []
    while inteiro:
        grupos.insert(0, inteiro[-3:])
        inteiro = inteiro[:-3]
    return sinal + ".".join(grupos or ["0"]) + (f",{fracao}" if fracao else "")


SITUACAO_CURTA = {
    "complete": "Concluído", "incomplete": "Pendente", "review": "Em revisão",
    "error": "Falha na consulta", "pending": "Não pesquisado",
    "searching": "Em pesquisa", "classifying": "Em pesquisa",
}


def situacao_curta(item: ItemView) -> str:
    return SITUACAO_CURTA.get(item.estado, item.situacao or NAO_INFORMADO)


# ---------------------------------------------------------------------------
# WordprocessingML mínimo
# ---------------------------------------------------------------------------
_W = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
      'relationships"')


def _rpr(*, pt: float, negrito=False, cor=TINTA, fonte=FONTE, italico=False,
         riscado=False, sublinhado=False) -> str:
    partes = [f'<w:rFonts w:ascii="{fonte}" w:hAnsi="{fonte}" '
              f'w:cs="{fonte}" w:eastAsia="{fonte}"/>']
    if negrito:
        partes.append("<w:b/><w:bCs/>")
    if italico:
        partes.append("<w:i/>")
    if riscado:
        partes.append("<w:strike/>")
    partes.append(f'<w:color w:val="{cor}"/>')
    meio = int(round(pt * 2))
    partes.append(f'<w:sz w:val="{meio}"/><w:szCs w:val="{meio}"/>')
    if sublinhado:
        partes.append('<w:u w:val="single"/>')
    return "<w:rPr>" + "".join(partes) + "</w:rPr>"


@dataclass(frozen=True)
class T:
    """Um trecho de texto com estilo; `link` vira hyperlink real."""
    texto: str
    negrito: bool = False
    cor: str = TINTA
    pt: float | None = None
    italico: bool = False
    riscado: bool = False
    link: str | None = None
    quebra_antes: bool = False


Trechos = str | T | list


def _como_trechos(conteudo: Trechos) -> list[T]:
    if conteudo is None:
        return []
    if isinstance(conteudo, str):
        return [T(conteudo)]
    if isinstance(conteudo, T):
        return [conteudo]
    saida: list[T] = []
    for parte in conteudo:
        saida.extend(_como_trechos(parte))
    return saida


class Documento:
    """Constrói o DOCX. Uma seção, retrato (compacto) ou paisagem."""

    def __init__(self, *, paisagem: bool):
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Cm, Pt

        self.doc = Document()
        secao = self.doc.sections[0]
        secao.page_height, secao.page_width = Cm(29.7), Cm(21.0)
        if paisagem:
            secao.orientation = WD_ORIENT.LANDSCAPE
            secao.page_width, secao.page_height = Cm(29.7), Cm(21.0)
        secao.left_margin = secao.right_margin = Cm(1.6)
        secao.top_margin = Cm(1.5)
        secao.bottom_margin = Cm(1.5)
        secao.header_distance = Cm(0.7)
        secao.footer_distance = Cm(0.6)
        self.secao = secao
        self.util_cm = (29.7 if paisagem else 21.0) - 3.2
        normal = self.doc.styles["Normal"]
        normal.font.name = FONTE
        normal.font.size = Pt(CORPO_PT)
        rpr = normal.element.get_or_add_rPr()
        rpr.get_or_add_rFonts().set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            "eastAsia", FONTE)
        self._fim = self.doc.element.body[-1]       # sectPr
        self._links: dict[str, str] = {}

    # -- infraestrutura -----------------------------------------------------
    def _rid(self, url: str) -> str:
        from docx.opc.constants import RELATIONSHIP_TYPE as RT

        if url not in self._links:
            self._links[url] = self.doc.part.relate_to(
                url, RT.HYPERLINK, is_external=True)
        return self._links[url]

    def _runs(self, conteudo: Trechos, *, pt: float, fonte: str = FONTE,
              cor: str = TINTA, negrito: bool = False) -> str:
        saida = []
        for t in _como_trechos(conteudo):
            texto = str(t.texto)
            if t.quebra_antes:
                saida.append("<w:r><w:br/></w:r>")
            props = _rpr(pt=t.pt or pt, negrito=t.negrito or negrito,
                         cor=(AZUL if t.link else
                              (t.cor if t.cor != TINTA else cor)),
                         fonte=fonte, italico=t.italico, riscado=t.riscado,
                         sublinhado=bool(t.link))
            linhas = texto.split("\n")
            corpo = "<w:br/>".join(
                f'<w:t xml:space="preserve">{escape(linha)}</w:t>'
                for linha in linhas)
            run = f"<w:r>{props}{corpo}</w:r>"
            if t.link:
                run = (f'<w:hyperlink r:id="{self._rid(t.link)}" '
                       f'w:history="1">{run}</w:hyperlink>')
            saida.append(run)
        return "".join(saida)

    @staticmethod
    def _ppr(*, alinhar="left", antes=0, depois=0, manter=False,
             entrelinha=1.12, quebra_pagina=False, recuo_cm=0.0,
             faixa: str | None = None, tab_direita_cm: float = 0.0) -> str:
        jc = {"left": "left", "right": "right", "center": "center",
              "justify": "both"}[alinhar]
        partes = []
        if manter:
            partes.append("<w:keepNext/>")
        if quebra_pagina:
            partes.append("<w:pageBreakBefore/>")
        if faixa:
            partes.append(f'<w:shd w:val="clear" w:color="auto" '
                          f'w:fill="{faixa}"/>')
        if tab_direita_cm:
            partes.append(f'<w:tabs><w:tab w:val="right" '
                          f'w:pos="{int(tab_direita_cm * CM)}"/></w:tabs>')
        partes.append(f'<w:spacing w:before="{int(antes * 20)}" '
                      f'w:after="{int(depois * 20)}" '
                      f'w:line="{int(240 * entrelinha)}" w:lineRule="auto"/>')
        if recuo_cm:
            partes.append(f'<w:ind w:left="{int(recuo_cm * CM)}" '
                          f'w:hanging="{int(0.35 * CM)}"/>')
        partes.append(f'<w:jc w:val="{jc}"/>')
        return "<w:pPr>" + "".join(partes) + "</w:pPr>"

    def _anexar(self, xml: str) -> None:
        from docx.oxml import parse_xml

        self._fim.addprevious(parse_xml(xml))

    # -- blocos ---------------------------------------------------------------
    def paragrafo(self, conteudo: Trechos, *, pt: float = CORPO_PT,
                  cor: str = TINTA, negrito=False, alinhar="left", antes=0,
                  depois=4, manter=False, quebra_pagina=False,
                  recuo_cm=0.0) -> None:
        self._anexar(
            f"<w:p {_W}>"
            + self._ppr(alinhar=alinhar, antes=antes, depois=depois,
                        manter=manter, quebra_pagina=quebra_pagina,
                        recuo_cm=recuo_cm)
            + self._runs(conteudo, pt=pt, cor=cor, negrito=negrito)
            + "</w:p>")

    def marcador(self, conteudo: Trechos, *, pt: float = CORPO_PT - 0.5,
                 cor: str = TINTA) -> None:
        self.paragrafo([T("•  ", cor=AZUL, negrito=True), *_como_trechos(
            conteudo)], pt=pt, cor=cor, depois=2, recuo_cm=0.45)

    def titulo_de_secao(self, texto: str, *, quebra_pagina=False) -> None:
        self.paragrafo(T(texto, negrito=True, cor=AZUL), pt=11.5, antes=10,
                       depois=4, manter=True, quebra_pagina=quebra_pagina)

    def tabela(self, colunas: list[Coluna], linhas: list[list[Trechos]], *,
               cabecalho: bool = True, pt: float = TABELA_PT,
               fundo_linhas: list[str | None] | None = None,
               fundo_cabecalho: str = CABECALHO_TABELA,
               bordas: bool = True, manter_junto: bool = False,
               fonte: str = FONTE_TABELA, negrito_coluna0: bool = False
               ) -> None:
        """
        Largura FIXA por coluna (soma = largura útil), cabeçalho repetido a
        cada página, linha que não se parte entre páginas.
        """
        larguras = [int(c.largura_cm * CM) for c in colunas]
        total = sum(larguras)
        linha_borda = (f'w:val="single" w:sz="4" w:space="0" '
                       f'w:color="{BORDA}"')
        nada = 'w:val="nil"'
        bordas_xml = (
            f"<w:tblBorders><w:top {linha_borda if bordas else nada}/>"
            f"<w:left {nada}/><w:bottom {linha_borda if bordas else nada}/>"
            f"<w:right {nada}/><w:insideH {linha_borda if bordas else nada}/>"
            f"<w:insideV {nada}/></w:tblBorders>")
        partes = [
            f"<w:tbl {_W}><w:tblPr>"
            f'<w:tblW w:w="{total}" w:type="dxa"/>'
            '<w:tblLayout w:type="fixed"/>' + bordas_xml +
            '<w:tblCellMar><w:top w:w="40" w:type="dxa"/>'
            '<w:left w:w="85" w:type="dxa"/><w:bottom w:w="40" w:type="dxa"/>'
            '<w:right w:w="85" w:type="dxa"/></w:tblCellMar>'
            '<w:tblLook w:val="04A0" w:firstRow="1" w:lastRow="0" '
            'w:firstColumn="0" w:lastColumn="0" w:noHBand="1" w:noVBand="1"/>'
            "</w:tblPr><w:tblGrid>"
            + "".join(f'<w:gridCol w:w="{w}"/>' for w in larguras)
            + "</w:tblGrid>"]

        def celula(conteudo, largura, alinhar, fundo, negrito, cor=TINTA,
                   manter=False):
            sombra = (f'<w:shd w:val="clear" w:color="auto" w:fill="{fundo}"/>'
                      if fundo else "")
            return (f'<w:tc><w:tcPr><w:tcW w:w="{largura}" w:type="dxa"/>'
                    f"{sombra}</w:tcPr><w:p>"
                    + self._ppr(alinhar=alinhar, entrelinha=1.05,
                                manter=manter)
                    + self._runs(conteudo, pt=pt, fonte=fonte, cor=cor,
                                 negrito=negrito)
                    + "</w:p></w:tc>")

        if cabecalho:
            partes.append("<w:tr><w:trPr><w:cantSplit/><w:tblHeader/>"
                          "</w:trPr>")
            for c, w in zip(colunas, larguras):
                partes.append(celula(T(c.titulo), w, c.alinhar,
                                     fundo_cabecalho, True, cor=AZUL,
                                     manter=True))
            partes.append("</w:tr>")
        for n, linha in enumerate(linhas):
            fundo = (fundo_linhas[n] if fundo_linhas and n < len(fundo_linhas)
                     else None)
            partes.append("<w:tr><w:trPr><w:cantSplit/></w:trPr>")
            for k, (c, w, valor) in enumerate(zip(colunas, larguras, linha)):
                partes.append(celula(valor, w, c.alinhar, fundo,
                                     negrito_coluna0 and k == 0,
                                     manter=manter_junto))
            partes.append("</w:tr>")
        partes.append("</w:tbl>")
        self._anexar("".join(partes))
        # parágrafo de respiro: o Word exige parágrafo entre tabelas
        self._anexar(f"<w:p {_W}>" + self._ppr(depois=2, entrelinha=0.6)
                     + "</w:p>")

    def caixa(self, conteudo: Trechos, *, fundo=ALERTA_FUNDO,
              cor=ALERTA_TINTA, pt: float = CORPO_PT - 0.5) -> None:
        """Bloco de destaque (pendência). Nunca só cor: o texto diz."""
        self.tabela([Coluna("", self.util_cm)], [[conteudo]], cabecalho=False,
                    pt=pt, fundo_linhas=[fundo], bordas=False, fonte=FONTE)

    def faixa_do_item(self, esquerda: str, direita: str) -> None:
        """
        PARÁGRAFO sombreado, e não tabela: o LibreOffice não respeita
        "manter com o próximo" em tabela, e a faixa ficava sozinha no pé
        da página (medido no processo de 210 itens).
        """
        self._anexar(
            f"<w:p {_W}>"
            + self._ppr(antes=8, depois=4, manter=True, faixa=AZUL,
                        tab_direita_cm=self.util_cm - 0.1)
            + self._runs([T(f" {esquerda}", negrito=True, cor=BRANCO)],
                         pt=10)
            + "<w:r><w:tab/></w:r>"
            + self._runs([T(f"{direita} ", negrito=True, cor=BRANCO)], pt=10)
            + "</w:p>")

    # -- cabeçalho e rodapé ---------------------------------------------------
    def cabecalho_e_rodape(self, view: CotacaoReportView, rotulo: str) -> None:
        from docx.oxml import parse_xml

        secao = self.secao
        secao.different_first_page_header_footer = True
        branding = dict(view.branding or {})
        imagem_topo = _imagem(branding, "cabecalho_img")
        imagem_base = _imagem(branding, "rodape_img")

        # primeira página: timbrado do órgão (imagem ou texto)
        primeiro = secao.first_page_header.paragraphs[0]
        if imagem_topo:
            self._figura(primeiro, imagem_topo, altura_max_cm=2.6)
        else:
            # Sem timbrado, o nome do órgão; sem órgão, nada — a abertura
            # já diz GovConnect, e repetir a marca no topo é ruído.
            orgao = view.orgao_emissor or str(branding.get("cabecalho") or "")
            if orgao:
                primeiro._p.append(parse_xml(
                    f"<w:r {_W}>" + _rpr(pt=8, negrito=True, cor=AZUL)
                    + f'<w:t xml:space="preserve">{escape(orgao.upper())}'
                    "</w:t></w:r>"))
                _linha_inferior(primeiro)
        # demais páginas: linha discreta
        demais = secao.header.paragraphs[0]
        demais._p.append(parse_xml(
            f"<w:r {_W}>" + _rpr(pt=7.5, cor=SECUNDARIO)
            + f'<w:t xml:space="preserve">{escape(rotulo)} · '
            f"{escape(view.nome)} · Revisão {view.revisao}</w:t></w:r>"))
        _linha_inferior(demais)

        texto_rodape = (f"GovConnect · Cotação {view.pesquisa_id[:8]} · "
                        f"Rev. {view.revisao} · Versão "
                        f"{view.hash_decisorio[:16]}… · "
                        f"{view.emitido_local}")
        for rodape in (secao.footer, secao.first_page_footer):
            par = rodape.paragraphs[0]
            if imagem_base:
                self._figura(par, imagem_base, altura_max_cm=1.4)
                par = rodape.add_paragraph()
            tab = int(self.util_cm * CM)
            par._p.insert(0, parse_xml(
                f"<w:pPr {_W}><w:pBdr><w:top w:val=\"single\" w:sz=\"4\" "
                f'w:space="4" w:color="{BORDA}"/></w:pBdr>'
                f'<w:tabs><w:tab w:val="right" w:pos="{tab}"/></w:tabs>'
                "<w:spacing w:before=\"0\" w:after=\"0\"/></w:pPr>"))
            props = _rpr(pt=7, cor=SECUNDARIO, fonte=FONTE_TABELA)
            par._p.append(parse_xml(
                f"<w:r {_W}>{props}<w:t xml:space=\"preserve\">"
                f"{escape(texto_rodape)}</w:t></w:r>"))
            par._p.append(parse_xml(f"<w:r {_W}>{props}<w:tab/>"
                                    "<w:t xml:space=\"preserve\">Página </w:t>"
                                    "</w:r>"))
            # Campo COMPLETO (begin/instr/separate/end), com a formatação
            # em cada parte: no `fldSimple` o LibreOffice ignora o rPr e o
            # número sai no corpo do texto, maior que o rodapé.
            for instrucao, depois in (("PAGE", " de "), ("NUMPAGES", "")):
                for parte in (
                        '<w:fldChar w:fldCharType="begin"/>',
                        f'<w:instrText xml:space="preserve"> {instrucao} '
                        "</w:instrText>",
                        '<w:fldChar w:fldCharType="separate"/>',
                        "<w:t>1</w:t>",
                        '<w:fldChar w:fldCharType="end"/>'):
                    par._p.append(parse_xml(f"<w:r {_W}>{props}{parte}</w:r>"))
                if depois:
                    par._p.append(parse_xml(
                        f"<w:r {_W}>{props}<w:t xml:space=\"preserve\">"
                        f"{depois}</w:t></w:r>"))

    def _figura(self, paragrafo, png: bytes, *, altura_max_cm: float) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm
        from PIL import Image

        with Image.open(io.BytesIO(png)) as img:
            largura, altura = img.size
        alvo_largura = self.util_cm
        if altura and alvo_largura * altura / largura > altura_max_cm:
            alvo_largura = altura_max_cm * largura / altura
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragrafo.add_run().add_picture(io.BytesIO(png),
                                        width=Cm(alvo_largura))

    def bytes(self) -> bytes:
        buffer = io.BytesIO()
        self.doc.save(buffer)
        return buffer.getvalue()


def _linha_inferior(paragrafo) -> None:
    from docx.oxml import parse_xml

    paragrafo._p.insert(0, parse_xml(
        f"<w:pPr {_W}><w:pBdr><w:bottom w:val=\"single\" w:sz=\"4\" "
        f'w:space="3" w:color="{BORDA}"/></w:pBdr></w:pPr>'))


def _imagem(branding: dict, chave: str) -> bytes | None:
    from ..branding import de_base64

    try:
        return de_base64(branding.get(chave)) if branding.get(chave) else None
    except Exception:  # noqa: BLE001 — timbrado ilegível não derruba o relatório
        return None


# ---------------------------------------------------------------------------
# Relatório de Cotação (compacto)
# ---------------------------------------------------------------------------
ROTULO_COMPACTO = "Relatório de Cotação de Preços"
ROTULO_ANALITICO = "Relatório Analítico Integral"


def _abertura(d: Documento, view: CotacaoReportView, titulo: str,
              subtitulo: str) -> None:
    d.paragrafo(T("GOVCONNECT · PESQUISA DE PREÇOS", negrito=True, cor=AZUL),
                pt=7.5, depois=2)
    d.paragrafo(T(titulo.upper(), negrito=True, cor=AZUL), pt=15, depois=2)
    d.paragrafo([T(view.nome, negrito=True), T(f"  ·  Revisão {view.revisao}"),
                 T(f"  ·  {subtitulo}", cor=SECUNDARIO)], pt=10, depois=8)

    coleta = view.coleta_local
    processo = (f"{view.processo_id[:8]} — {view.processo_objeto}"
                if view.processo_id and view.processo_objeto else
                view.processo_id[:8] if view.processo_id else
                "Pesquisa autônoma (não vinculada a processo)")
    pares = [
        ("Órgão", view.orgao_emissor or NAO_INFORMADO,
         "Data-base", view.data_base),
        ("Processo", processo, "Coleta", coleta),
        ("Responsável", view.responsavel, "Situação da pesquisa",
         view.situacao),
        ("Perfil normativo", f"{view.perfil_nome} ({view.perfil_base_legal})",
         "Método", ", ".join(view.metodos_aplicados) or "Não aplicado"),
        ("Local de referência", view.local, "Emitido em",
         view.emitido_local),
    ]
    rotulo = 3.1
    valor = (d.util_cm - 2 * rotulo) / 2
    d.tabela(
        [Coluna("", rotulo), Coluna("", valor), Coluna("", rotulo),
         Coluna("", valor)],
        [[T(a, cor=SECUNDARIO), T(b, negrito=True), T(c, cor=SECUNDARIO),
          T(e, negrito=True)] for a, b, c, e in pares],
        cabecalho=False, fundo_linhas=[FAIXA] * len(pares), pt=8)
    d.paragrafo([T("Objeto: ", negrito=True), T(view.objeto)], pt=8.5,
                depois=3, alinhar="justify")
    if view.motivo_da_revisao:
        d.paragrafo([T("Motivo da revisão: ", negrito=True),
                     T(view.motivo_da_revisao)], pt=NOTA_PT, cor=SECUNDARIO)


def _aviso_de_pendencias(d: Documento, view: CotacaoReportView) -> None:
    pendentes = [i for i in view.itens if not i.concluido]
    if not pendentes:
        return
    lista = "; ".join(f"{i.numero:02d} ({situacao_curta(i).lower()}, "
                      f"{i.consideradas}/{i.minimo})" for i in pendentes[:12])
    if len(pendentes) > 12:
        lista += f"; e mais {len(pendentes) - 12}"
    d.caixa([T("ATENÇÃO — ", negrito=True),
             T(f"{len(pendentes)} de {len(view.itens)} itens não concluídos. "
               "O valor global da contratação ainda não está consolidado; o "
               "subtotal abaixo soma apenas os itens concluídos. "),
             T("Itens: ", negrito=True), T(lista + ".")])


def _quadro_resumo(d: Documento, view: CotacaoReportView) -> None:
    d.titulo_de_secao("Resumo da cotação")
    if not view.itens:
        d.paragrafo("Esta cotação não tem itens. Não há preço a relatar.",
                    cor=SECUNDARIO)
        return
    colunas = [Coluna("Nº", 0.9, "center"), Coluna("Item", 6.0),
               Coluna("Qtd. / Unid.", 2.5, "right"),
               Coluna("Preços usados", 2.8, "center"),
               Coluna("Unitário", 2.2, "right"), Coluna("Total", 2.5, "right"),
               Coluna("Situação", 2.2)]
    _ajustar(colunas, d.util_cm, elastica=1)
    linhas, fundos = [], []
    for item in view.itens:
        concluido = item.concluido
        linhas.append([
            f"{item.numero:02d}",
            item.titulo,
            f"{quantidade(item.quantidade)} {item.unidade}",
            f"{item.consideradas}",
            (moeda(item.preco_unitario) + (" *" if item.ajuste_teto else ""))
            if concluido else "—",
            T(moeda(item.preco_total), negrito=True) if concluido else "—",
            T(situacao_curta(item), negrito=not concluido,
              cor=ALERTA_TINTA if not concluido else TINTA),
        ])
        fundos.append(None if concluido else ALERTA_FUNDO)
    d.tabela(colunas, linhas, fundo_linhas=fundos)

    concluidos = f"{view.concluidos} de {len(view.itens)}"
    # Nenhum concluído: "R$ 0,00" leria como preço zero — é ausência.
    subtotal = (moeda(view.subtotal_concluidos) if view.concluidos
                else "Nenhum item concluído")
    totais = [[T(f"Subtotal dos itens concluídos ({concluidos})",
                 negrito=True), T(subtotal, negrito=True)]]
    if view.valor_global is not None:
        totais.append([T("Valor global estimado", negrito=True, cor=AZUL),
                       T(moeda(view.valor_global), negrito=True, cor=AZUL)])
    d.tabela([Coluna("", d.util_cm - 4.2, "right"),
              Coluna("", 4.2, "right")], totais, cabecalho=False,
             fundo_linhas=[FAIXA] * len(totais), pt=9, fonte=FONTE)
    if view.valor_global is None:
        d.paragrafo(
            f"{view.nao_concluidos} item(ns) não concluído(s): valor global "
            "não consolidado. Itens pendentes não entram no subtotal e não "
            "têm preço estimado neste quadro.", pt=NOTA_PT,
            cor=ALERTA_TINTA)
    if any(i.ajuste_teto for i in view.itens):
        d.paragrafo(
            "* Valor limitado à mediana dos preços: a estimativa se apoia só "
            "em sistema oficial de preços e o perfil normativo da pesquisa "
            "manda que ela não supere a mediana. O valor calculado pelo "
            "método e o adotado estão no detalhamento do item.",
            pt=NOTA_PT, cor=SECUNDARIO)
    d.paragrafo(
        f"Mínimo por item: {view.minimo} preços. “Preços usados” conta as "
        "referências que compõem o cálculo. Valores em reais com duas casas "
        "(arredondamento meio para cima); total do item = unitário × "
        "quantidade.", pt=NOTA_PT, cor=SECUNDARIO)


def _metodo(d: Documento, view: CotacaoReportView) -> None:
    d.titulo_de_secao("Método e critérios")
    linhas = [
        ["Perfil normativo", f"{view.perfil_nome} — {view.perfil_base_legal}"],
        ["Regra de formação", view.regra_rotulo],
        ["Mínimo por item", f"{view.minimo} referências consideradas"
         + ("; admite concluir com menos mediante justificativa registrada"
            if view.admite_menos else "")],
    ]
    busca = (f"últimos {view.janela_dias} dias contados da data-base"
             if view.janela_dias else "janela temporal não informada")
    busca += (f"; recorte geográfico: {view.recorte_uf}" if view.recorte_uf
              else "; sem recorte geográfico")
    linhas.append(["Critérios de busca", busca + "; código de catálogo "
                   "(CATMAT/CATSER) usado quando informado, não exigido"])
    # Revisão 02 (05_RELATORIO §3): metodologia sucinta e só do que de
    # fato ocorreu. A amostra é a USADA; a triagem só é descrita quando
    # agiu nesta pesquisa (versão lida dos motivos gravados).
    linhas.append([
        "Seleção da amostra",
        "Preços praticados em contratações públicas, na unidade do item, "
        "escolhidos por prioridade da fonte e comparabilidade com a "
        "especificação — nunca pelo menor preço. Na cotação automática, até "
        "15 preços por item."])
    if view.triagem_versoes:
        linhas.append([
            "Triagem antes da cesta",
            "Ficam fora da série: valor inválido, unidade sem conversão "
            "comprovada pela fonte, a mesma linha de contratação repetida e "
            "preço atípico corroborado por pelo menos 8 preços comparáveis "
            "de origens independentes. Com amostra menor, o preço atípico "
            "fica no cálculo, sinalizado para conferência. Política "
            + ", ".join(view.triagem_versoes) + "."])
    linhas.append([
        "Preços atípicos",
        "Sinalizados pela distância da mediana do item (desvio absoluto "
        "mediano e intervalo interquartil). A sinalização não retira o "
        "preço: o servidor remove ou mantém, e a decisão consta do item."])
    if view.fundamento_desconsideracao:
        linhas.append([
            "Desconsideração de valores",
            f"Fundamento: {view.fundamento_desconsideracao}. As "
            "desconsiderações que afetaram a série de cada item estão no "
            "detalhamento, com o critério."])
    d.tabela([Coluna("", 4.2), Coluna("", d.util_cm - 4.2)],
             [[T(a, negrito=True, cor=SECUNDARIO), b] for a, b in linhas],
             cabecalho=False, pt=8)
    for observacao in view.perfil_observacoes:
        d.paragrafo(observacao, pt=NOTA_PT, cor=SECUNDARIO)


def _celula_fonte(ref: ReferenciaView) -> list[T]:
    partes = [T(ref.fonte, negrito=True)]
    if ref.inciso:
        partes.append(T(f" · inc. {ref.inciso}", cor=SECUNDARIO))
    partes.append(T(ref.id_externo, link=ref.link, quebra_antes=True,
                    cor=SECUNDARIO, pt=7.5))
    return partes


def _celula_orgao(ref: ReferenciaView) -> list[T]:
    partes = [T(ref.orgao)]
    if ref.uf:
        partes.append(T(f" · {ref.uf}", cor=SECUNDARIO))
    return partes


def _celula_origem(ref: ReferenciaView) -> list[T]:
    """
    Órgão e o que a fonte disse da contratação — só o que existe no
    registro (Revisão 02, 05 §2: nunca preencher com placeholder).
    """
    partes = _celula_orgao(ref)
    detalhes = []
    if ref.modalidade:
        detalhes.append(ref.modalidade)
    if ref.item_da_compra:
        detalhes.append(f"item {ref.item_da_compra}")
    if detalhes:
        partes.append(T(" · ".join(detalhes), quebra_antes=True,
                        cor=SECUNDARIO, pt=7))
    if ref.fornecedor:
        fornecedor = ref.fornecedor + (f" ({ref.fornecedor_cnpj})"
                                       if ref.fornecedor_cnpj else "")
        partes.append(T(f"Fornecedor: {fornecedor}", quebra_antes=True,
                        cor=SECUNDARIO, pt=7))
    return partes


def _celula_preco(ref: ReferenciaView) -> list[T]:
    partes = [T(moeda(ref.valor) + (" ⚑" if ref.sinalizada else ""),
                negrito=True)]
    if ref.valor_original is not None and ref.valor_original != ref.valor:
        partes.append(T(f"orig. {moeda(ref.valor_original)} / "
                        f"{ref.unidade_original}", quebra_antes=True,
                        cor=SECUNDARIO, pt=7))
    return partes


def _celula_contratacao(ref: ReferenciaView) -> list[T]:
    """
    O link da PÁGINA DA CONTRATAÇÃO, sobre o número oficial impresso.
    Sem página pública resolvida, diz isso — nunca aponta o JSON da API.
    """
    if ref.contratacao_url and ref.identificacao:
        return [T(ref.identificacao, link=ref.contratacao_url, pt=7,
                  cor=AZUL),
                T(ref.contratacao_rotulo, quebra_antes=True,
                  cor=SECUNDARIO, pt=6.5)]
    return [T("Não disponível na fonte", cor=SECUNDARIO, pt=7)]


def _memoria_para_imprimir(item: ItemView) -> list[str]:
    """
    A memória gravada, sem contradizer a situação do item (Revisão 02,
    D06/AC-32): concluído não sai "PROVISÓRIO" nem "aguardam decisão" —
    a conclusão do revisor já está na própria memória.
    """
    if not item.concluido:
        return list(item.memoria)
    descartar = ("valor PROVISÓRIO", "Análise crítica:")
    return [linha for linha in item.memoria
            if not linha.startswith(descartar)
            and "aguardam decisão" not in linha]


def _detalhe_do_item(d: Documento, view: CotacaoReportView,
                     item: ItemView) -> None:
    d.faixa_do_item(f"ITEM {item.numero:02d} · {item.titulo.upper()}",
                    situacao_curta(item).upper())
    # faixa + especificação + dados + indicadores andam JUNTOS: título
    # separado do seu bloco é o defeito de paginação que o pacote proíbe
    d.paragrafo([T("Especificação: ", negrito=True), T(item.descricao)],
                pt=8.5, depois=3, alinhar="justify", manter=True)
    dados = [T("Unidade: ", cor=SECUNDARIO), T(item.unidade, negrito=True),
             T("   Quantidade: ", cor=SECUNDARIO),
             T(quantidade(item.quantidade), negrito=True)]
    # Revisão 02 (D11): só código de catálogo REAL; o número da linha da
    # planilha não é catálogo ("Catálogo: Catálogo 1").
    if item.codigo and item.tipo_catalogo in ("CATMAT", "CATSER"):
        dados += [T(f"   {item.tipo_catalogo}: ", cor=SECUNDARIO),
                  T(item.codigo)]
    d.paragrafo(dados, pt=8.5, depois=4, manter=True)

    if item.concluido:
        unitario = moeda(item.preco_unitario)
        total = moeda(item.preco_total)
    elif item.preco_unitario is not None:
        unitario = f"{moeda(item.preco_unitario)} (em revisão)"
        total = "Não consolidado"
    else:
        unitario, total = "Não estimado", "Não estimado"
    kpis = [("Método", item.metodo_rotulo or "Não aplicado"),
            ("Preços usados", f"{item.consideradas} (mínimo {item.minimo})"),
            ("Valor unitário", unitario), ("Valor do item", total)]
    largura = d.util_cm / 4
    # UMA linha (rótulo e valor na mesma célula): linha não se parte
    # entre páginas, então rótulo e valor nunca se separam.
    d.tabela([Coluna("", largura)] * 4,
             [[[T(r, cor=SECUNDARIO, pt=7.5),
                T(v, negrito=True, pt=10, quebra_antes=True)]
               for r, v in kpis]],
             cabecalho=False, fundo_linhas=[FAIXA], bordas=False,
             fonte=FONTE)
    if item.ajuste_teto:
        pelo_metodo, adotado = item.ajuste_teto
        d.paragrafo(
            [T("Valor adotado: ", negrito=True),
             T(f"{item.metodo_rotulo or 'método'} dos preços = "
               f"{moeda(pelo_metodo)}; adotada a mediana, {moeda(adotado)}, "
               "porque a estimativa se apoia só em sistema oficial de preços "
               f"e o perfil ({view.perfil_nome}) limita o valor à mediana.")],
            pt=8, cor=SECUNDARIO, depois=3)

    if item.deficiencias:
        texto = "; ".join(item.deficiencias)
        d.caixa([T("PENDENTE — " if not item.concluido else "RESSALVA — ",
                   negrito=True),
                 T(f"{item.consideradas} preço(s) usado(s), mínimo "
                   f"{item.minimo}. "),
                 T(texto[0].upper() + texto[1:] + "."),
                 T(" Este item não integra o subtotal." if not item.concluido
                   else "")])
    for problema in item.integridade:
        d.caixa([T("Atenção à cesta — ", negrito=True), T(problema + ".")])

    e = item.estatisticas
    if e.get("quantidade"):
        partes = [T(f"Série usada ({e.get('quantidade')}): ", negrito=True)]
        for rotulo, chave in (("menor", "menor"), ("média", "media"),
                              ("mediana", "mediana"), ("maior", "maior")):
            if e.get(chave) is not None:
                partes.append(T(f"{rotulo} {moeda(e[chave])}   "))
        if e.get("coeficiente_variacao") is not None:
            partes.append(T(f"coef. de variação "
                            f"{pct(e['coeficiente_variacao'])}"))
        d.paragrafo(partes, pt=8, cor=SECUNDARIO, depois=3)

    # Revisão 02 (05 §2): TODAS as referências usadas, com a contratação
    if item.consideradas:
        exibidas = len(item.exibidas)
        cabeca = (f"Preços usados ({item.consideradas})" if exibidas >=
                  item.consideradas else
                  f"Preços usados — exibidos {exibidas} de "
                  f"{item.consideradas} (menor, quartis, mediana e maior)")
        d.paragrafo(T(cabeca, negrito=True), pt=8.5, depois=2, manter=True)
        colunas = [Coluna("Nº", 0.8, "right"), Coluna("Data", 2.1),
                   Coluna("Órgão · contratação", 5.4),
                   Coluna("Fonte · natureza", 3.0),
                   Coluna("Unitário", 2.3, "right"),
                   Coluna("Contratação", 4.4)]
        _ajustar(colunas, d.util_cm, elastica=2)
        d.tabela(colunas, [[
            str(n), r.data, _celula_origem(r),
            [T(r.fonte, negrito=True, pt=7.5),
             T((f"inc. {r.inciso} · " if r.inciso else "") + r.natureza,
               quebra_antes=True, cor=SECUNDARIO, pt=7)],
            _celula_preco(r), _celula_contratacao(r),
        ] for n, r in enumerate(item.exibidas, start=1)], pt=7.5)
        if exibidas < item.consideradas:
            d.paragrafo(f"Todas as {item.consideradas} referências compõem o "
                        "cálculo e estão na Memória Analítica (XLSX).",
                        pt=7.5, cor=SECUNDARIO, depois=3)
        if any(r.sinalizada for r in item.exibidas):
            d.paragrafo("⚑ preço atípico que está no cálculo: veja abaixo a "
                        "decisão sobre ele.", pt=7.5, cor=SECUNDARIO)

    # Revisão 02 (08_COMPLIANCE, casos B e C): só o que AFETOU a série,
    # com a autoria certa; o inventário da busca fica na memória analítica
    saidas = item.saidas_da_serie
    if saidas:
        d.paragrafo(T(f"Desconsiderados da série ({len(saidas)})",
                      negrito=True), pt=8.5, depois=2, manter=True)
        linhas_saida = []
        for r in saidas[:10]:
            motivo = _motivo_curto(r)
            celula = [T(r.autoria_da_saida[:1].upper()
                        + r.autoria_da_saida[1:], negrito=True)]
            if motivo:
                celula.append(T(motivo, quebra_antes=True, cor=SECUNDARIO,
                                pt=7))
            linhas_saida.append([T(moeda(r.valor), riscado=True, cor=EXCLUIDA),
                                 _celula_orgao(r), celula])
        d.tabela([Coluna("Preço", 2.4, "right"), Coluna("Órgão · UF", 5.0),
                  Coluna("Como saiu", d.util_cm - 7.4)], linhas_saida, pt=7.5)
        if len(saidas) > 10:
            d.paragrafo(f"E mais {len(saidas) - 10}, na Memória Analítica.",
                        pt=7.5, cor=SECUNDARIO)
    if item.nao_usadas > len(saidas):
        d.paragrafo(
            f"Outros {item.nao_usadas - len(saidas)} resultado(s) da busca "
            "não entraram na série (unidade sem conversão comprovada, "
            "repetição da mesma contratação, comparabilidade baixa ou "
            "reserva além do limite). Constam da Memória Analítica.",
            pt=7.5, cor=SECUNDARIO, depois=3)

    # decisões sobre preços atípicos que ficaram no cálculo
    for ref, decisao in item.mantidas:
        d.marcador([T("Preço atípico mantido: ", negrito=True),
                    T(f"{moeda(ref.valor)} — {ref.orgao}"
                      + (f"/{ref.uf}" if ref.uf else "")
                      + f", {ref.data}. Mantido no cálculo por decisão do "
                        "servidor depois do alerta, em "
                        f"{_data_curta(decisao.get('registrado_em'))}.")],
                   pt=8)
    mantidas_ids = {r.id for r, _ in item.mantidas}
    if not item.concluido:
        for ref in item.considerada_lista:
            if ref.sinalizada and ref.id not in mantidas_ids:
                d.marcador([T("Preço a conferir: ", negrito=True),
                            T(f"{moeda(ref.valor)} — {ref.orgao}"
                              + (f"/{ref.uf}" if ref.uf else "")
                              + f", {ref.data}.")], pt=8)

    memoria = _memoria_para_imprimir(item)
    if memoria:
        d.paragrafo(T("Memória de cálculo", negrito=True), pt=8.5, depois=1,
                    manter=True, antes=2)
        for linha in memoria:
            d.marcador(linha, pt=8)
    if item.excecao:
        x = item.excecao
        d.paragrafo([T("Exceção à regra dos três: ", negrito=True),
                     T(f"{x.get('fundamento_texto') or x.get('fundamento') or NAO_INFORMADO} "
                       f"— {x.get('base_normativa') or NAO_INFORMADO}. "),
                     T(f"Justificativa: {x.get('texto') or NAO_INFORMADO}. "),
                     T(f"Aprovação: {x.get('aprovador') or NAO_INFORMADO} "
                       "(declarada como juntada aos autos).")], pt=8)
    for motivo in item.excecoes_invalidadas:
        d.marcador(f"Exceção anterior invalidada: {motivo}.", pt=8)
    if item.metodo_do_revisor:
        m = item.metodo_do_revisor
        d.paragrafo([T("Método definido pelo revisor: ", negrito=True),
                     T(f"{m.get('metodo')} — {m.get('justificativa')}")],
                    pt=8)
    if not item.concluido:
        for pendencia in item.analise_critica:
            d.marcador([T("Análise crítica pendente: ", negrito=True),
                        T(pendencia)], pt=8)
    for marca in item.marcas_percentuais[:3]:
        rotulo = ("elevado" if marca.get("marca") == "elevado_sinalizado"
                  else "inexequível")
        d.marcador(f"{moeda(Decimal(str(marca.get('valor'))))} sinalizado "
                   f"como {rotulo} pelo critério percentual da cotação; não "
                   "excluído automaticamente.", pt=8)
    if item.cotacoes_diretas:
        d.paragrafo(T(f"Cotações diretas com fornecedores "
                      f"({len(item.cotacoes_diretas)}) — Lei 14.133/2021, "
                      "art. 23, § 1º, IV", negrito=True), pt=8.5, manter=True)
        d.tabela(
            [Coluna("Fornecedor", 5.0), Coluna("CNPJ", 3.2),
             Coluna("Data", 2.0), Coluna("Unitário", 2.4, "right"),
             Coluna("Comprovante (SHA-256)", d.util_cm - 12.6)],
            [[str(c.get("fornecedor") or NAO_INFORMADO),
              str(c.get("cnpj") or NAO_INFORMADO),
              str(c.get("data") or NAO_INFORMADO)[:10],
              moeda(Decimal(str(c["valor"])) if c.get("valor") else None),
              T(str(c.get("comprovante") or NAO_INFORMADO), pt=7)]
             for c in item.cotacoes_diretas])


_CODIGOS = re.compile(r"\s*\[[A-Z_, ]+;\s*política [^\]]+\]")
_ROTULOS = re.compile(r"^(desconsiderada pelo Jev antes da cesta|requer "
                      r"revisão \(triagem pré-cesta\)):\s*")


def _motivo_curto(ref: ReferenciaView) -> str:
    """O critério, sem códigos internos nem versão de política."""
    texto = ref.motivo.split("; ")[-1] if ref.motivo else ""
    if "Jev" in ref.motivo or "triagem" in ref.motivo:
        texto = next((m for m in ref.motivo.split("; ")
                      if m.startswith(("desconsiderada pelo Jev",
                                       "requer revisão"))), texto)
    texto = _ROTULOS.sub("", _CODIGOS.sub("", texto)).strip()
    if texto == "retirada da cesta pelo usuário":
        return ""
    return texto[:240]


def _data_curta(valor) -> str:
    texto = str(valor or "")
    if len(texto) >= 10 and texto[4] == "-":
        return f"{texto[8:10]}/{texto[5:7]}/{texto[:4]}"
    return texto or NAO_INFORMADO


def _fontes(d: Documento, view: CotacaoReportView) -> None:
    d.titulo_de_secao("Fontes consultadas")
    if not view.fontes:
        d.paragrafo("Nenhuma fonte foi consultada nesta cotação — não há "
                    "origem de preço a relatar.", cor=SECUNDARIO)
        return
    colunas = [Coluna("Fonte", 4.4), Coluna("Enquadramento", 4.6),
               Coluna("Refs.", 1.3, "right"), Coluna("Consid.", 1.5, "right"),
               Coluna("Coleta (Brasília)", 3.0)]
    _ajustar(colunas, d.util_cm, elastica=1)
    linhas = []
    for f in view.fontes:
        desfechos = "; ".join(f"{r}: {n}" for r, n in f.desfechos.items())
        linhas.append([
            [T(f.nome, negrito=True)] + ([T(f"Desfechos por item — "
                                             f"{desfechos}", quebra_antes=True,
                                             cor=SECUNDARIO, pt=7.5)]
                                          if desfechos else []),
            f.enquadramento, str(f.referencias), str(f.consideradas),
            f"{_local(f.primeira_coleta)}\n{_local(f.ultima_coleta)}"])
    d.tabela(colunas, linhas)


def _alertas(d: Documento, view: CotacaoReportView) -> None:
    from .relatorio import texto_da_ciencia

    d.titulo_de_secao("Alertas e ressalvas")
    # Revisão 02 (D12): ocorrência operacional de fonte e a frase de
    # política com código ficam na trilha e na memória analítica; aqui só
    # o que o leitor precisa para usar o preço.
    # o "candidatos discrepantes" genérico também sai: cada item já diz
    # qual preço é atípico, de onde veio e o que foi decidido sobre ele
    alertas = [(t, x) for t, x in view.alertas
               if not t.lower().startswith(("ocorrência",
                                            "candidatos discrepantes"))]
    for titulo, texto in alertas:
        partes = [T(f"{titulo}: ", negrito=True)] if titulo else []
        d.marcador(partes + [T(texto)], pt=8.5)
    if view.ciencia:
        c = dict(view.ciencia)
        quantos = len(c.get("alertas") or [])
        d.marcador([T("Conferência antes da emissão: ", negrito=True),
                    T(f"{quantos} preço(s) sinalizado(s) foram apresentados "
                      "antes de emitir e mantidos como estavam; a ciência "
                      "ficou registrada para esta versão do resultado. "
                      "Ciência não é aprovação administrativa do preço.")],
                   pt=8.5)
    if not alertas and not view.ciencia:
        d.marcador("Nenhuma ressalva para esta versão do resultado.", pt=8.5)


def _identificacao(d: Documento, view: CotacaoReportView, *,
                   anexos: str, hash_completo: bool = False) -> None:
    d.titulo_de_secao("Identificação do documento")
    linhas = [
        ["Pesquisa", view.pesquisa_id],
        ["Revisão", str(view.revisao)],
        ["Identificador da versão",
         [T(view.hash_decisorio, pt=7.5, negrito=True)] if hash_completo
         else [T(view.hash_decisorio[:16] + "…", negrito=True),
               T(" (SHA-256 do conteúdo decisório; completo no Relatório "
                 "Analítico e na Memória Analítica)", cor=SECUNDARIO, pt=7)]],
        ["Conferência", "No GovConnect, em Pesquisa de Preços → Conferir "
         "identificador. Este documento não tem certificação digital nem "
         "código de autenticação externo."],
        ["Emitido em", view.emitido_em_texto],
        ["Pesquisa criada em", view.criada_em],
        ["Coleta realizada", view.coleta_local],
        ["Referências", f"{view.total_referencias} coletadas, "
         f"{view.total_consideradas} consideradas no cálculo"],
        ["Documentos da mesma versão", anexos],
    ]
    d.tabela([Coluna("", 4.2), Coluna("", d.util_cm - 4.2)],
             [[T(a, negrito=True, cor=SECUNDARIO), b] for a, b in linhas],
             cabecalho=False, pt=8)


def _local(texto_utc: str) -> str:
    """ "17/10/2026 18:58" em UTC → horário de Brasília; inalterado se não der."""
    from datetime import datetime, timezone

    from .relatorio_view import BRASILIA

    try:
        momento = datetime.strptime(str(texto_utc), "%d/%m/%Y %H:%M")
    except ValueError:
        return str(texto_utc)
    return momento.replace(tzinfo=timezone.utc).astimezone(
        BRASILIA).strftime("%d/%m/%Y %H:%M")


def _ajustar(colunas: list[Coluna], util_cm: float, *, elastica: int) -> None:
    """A coluna `elastica` absorve a diferença até a largura útil."""
    soma = sum(c.largura_cm for c in colunas)
    alvo = colunas[elastica]
    colunas[elastica] = Coluna(alvo.titulo,
                               round(alvo.largura_cm + util_cm - soma, 2),
                               alvo.alinhar)


def relatorio_de_cotacao_docx(view: CotacaoReportView) -> bytes:
    d = Documento(paisagem=False)
    d.cabecalho_e_rodape(view, ROTULO_COMPACTO)
    _abertura(d, view, ROTULO_COMPACTO, f"Data-base {view.data_base}")
    _aviso_de_pendencias(d, view)
    _quadro_resumo(d, view)
    _metodo(d, view)
    if view.itens:
        d.titulo_de_secao("Detalhamento dos itens", quebra_pagina=True)
        for item in view.itens:
            _detalhe_do_item(d, view, item)
    _fontes(d, view)
    _alertas(d, view)
    _identificacao(d, view, anexos=(
        "Memória Analítica (XLSX) e Relatório Analítico Integral (PDF): "
        "todos os resultados da busca, usados ou não, com o mesmo "
        "identificador da versão."))
    return d.bytes()


# ---------------------------------------------------------------------------
# Relatório Analítico Integral (paisagem)
# ---------------------------------------------------------------------------
def _linhas_do_anexo(item: ItemView) -> tuple[list[list[Trechos]], list]:
    linhas, fundos = [], []
    for n, r in enumerate(item.referencias, start=1):
        preco = [T(moeda(r.valor), negrito=r.considerada,
                   riscado=r.status == "rejected",
                   cor=EXCLUIDA if r.status == "rejected" else TINTA)]
        if r.valor_original is not None and r.valor_original != r.valor:
            preco.append(T(f"orig. {moeda(r.valor_original)} / "
                           f"{r.unidade_original}", quebra_antes=True,
                           cor=SECUNDARIO, pt=7))
        situacao = [T(r.situacao, negrito=True,
                      cor=TINTA if r.considerada else
                      (EXCLUIDA if r.status == "rejected" else SECUNDARIO))]
        linhas.append([
            str(n), situacao,
            [T(r.fonte, negrito=True),
             T(f"inc. {r.inciso} · {r.natureza}" if r.inciso else r.natureza,
               quebra_antes=True, cor=SECUNDARIO, pt=7)],
            [T(r.orgao), T(" · ".join(x for x in (r.uf, r.data) if x),
                           quebra_antes=True, cor=SECUNDARIO, pt=7)],
            f"{quantidade(r.quantidade)} {r.unidade}",
            preco,
            pct(r.compatibilidade, 0) if r.compatibilidade is not None
            else "—",
            [T(r.id_externo, link=r.contratacao_url or r.dados_url or r.link,
               pt=7),
             T(f"evid. {r.evidencia[:12]}…" if r.evidencia else "",
               quebra_antes=bool(r.evidencia), cor=SECUNDARIO, pt=6.5)],
            T(r.motivo or ("—" if r.considerada else ""), pt=7.5),
        ])
        fundos.append(None if r.considerada else "F7F8FA")
    return linhas, fundos


def relatorio_analitico_docx(view: CotacaoReportView) -> bytes:
    d = Documento(paisagem=True)
    d.cabecalho_e_rodape(view, ROTULO_ANALITICO)
    _abertura(d, view, ROTULO_ANALITICO,
              f"{view.total_referencias} referências")
    d.paragrafo(
        "Todas as referências coletadas para cada item — consideradas no "
        "cálculo e não consideradas —, com situação, motivo, link de "
        "verificação quando a fonte o fornece e a impressão digital da "
        "evidência. Mesmo identificador da versão do Relatório de Cotação.",
        pt=8.5, cor=SECUNDARIO)
    colunas = [Coluna("Nº", 0.9, "right"), Coluna("Situação", 2.5),
               Coluna("Fonte · inciso · natureza", 4.4),
               Coluna("Órgão · UF · data", 4.6),
               Coluna("Qtd. · unid.", 2.1, "right"),
               Coluna("Preço unit.", 2.3, "right"),
               Coluna("Compat.", 1.7, "right"),
               # cabe o controle PNCP inteiro ("…-1-000006/2026") numa
               # linha: identificador partido perde o link em leitores
               Coluna("Identificador", 4.2),
               Coluna("Motivo", 3.8)]
    _ajustar(colunas, d.util_cm, elastica=8)
    for item in view.itens:
        d.faixa_do_item(f"ITEM {item.numero:02d} · {item.titulo.upper()}",
                        situacao_curta(item).upper())
        unitario = (moeda(item.preco_unitario) if item.concluido else
                    f"{moeda(item.preco_unitario)} (provisório)"
                    if item.preco_unitario is not None else "Não estimado")
        d.paragrafo([T("Método: ", cor=SECUNDARIO),
                     T(item.metodo_rotulo or "Não aplicado", negrito=True),
                     T("   Válidas: ", cor=SECUNDARIO),
                     T(f"{item.consideradas}/{item.minimo}", negrito=True),
                     T("   Unitário: ", cor=SECUNDARIO),
                     T(unitario, negrito=True),
                     T("   Referências: ", cor=SECUNDARIO),
                     T(f"{item.total_referencias} "
                       f"({item.consideradas} consideradas, "
                       f"{item.total_referencias - item.consideradas} não)")],
                    pt=8.5, depois=3, manter=True)
        if not item.referencias:
            d.paragrafo("Nenhuma referência coletada para este item.",
                        pt=8.5, cor=SECUNDARIO)
            continue
        linhas, fundos = _linhas_do_anexo(item)
        d.tabela(colunas, linhas, fundo_linhas=fundos, pt=7.5)
        for anomalia in item.anomalias:
            d.marcador(f"Discrepante sinalizado: "
                       f"{moeda(Decimal(str(anomalia.get('valor'))))} — "
                       f"{anomalia.get('motivo')} (critério "
                       f"{anomalia.get('criterio')}); não excluído "
                       "automaticamente.", pt=7.5)
    _fontes(d, view)
    _identificacao(d, view, anexos=(
        "Relatório de Cotação de Preços (PDF/DOCX) e Memória Analítica "
        "(XLSX), com o mesmo identificador da versão."), hash_completo=True)
    return d.bytes()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def links_da_vista(view: CotacaoReportView, *,
                   analitico: bool = False) -> dict[str, str]:
    """
    Texto impresso → URL real. Relatório de cotação: o número oficial da
    contratação → a PÁGINA pública dela (Revisão 02, D08). Analítico: o
    identificador da referência → a página da contratação, ou, sem ela, o
    endpoint de dados da fonte (é o anexo técnico).
    """
    if not analitico:
        return {r.identificacao: r.contratacao_url
                for i in view.itens for r in i.referencias
                if r.contratacao_url and r.identificacao}
    return {r.id_externo: (r.contratacao_url or r.dados_url or r.link)
            for i in view.itens for r in i.referencias
            if (r.contratacao_url or r.dados_url or r.link)
            and r.id_externo and r.id_externo != NAO_INFORMADO}


def carimbar_links(pdf: bytes, links: dict[str, str]) -> bytes:
    """
    Grava o hyperlink sobre cada identificador impresso, no PDF.

    POR QUE EXISTE: o LibreOffice 7.4 da produção descarta os hyperlinks
    do DOCX ao exportar (medido no contêiner, 16/10/2026 — e é por isso que
    o relatório real de 907 páginas tinha zero links); o 26.2 local os
    preserva. Em vez de depender da versão do conversor, a área clicável
    é posta sobre a palavra exata, com a URL que a projeção já tem.

    Só age quando o conversor não exportou link nenhum — nunca duplica.
    """
    if not links:
        return pdf
    import pymupdf

    doc = pymupdf.open(stream=pdf, filetype="pdf")
    if any(link.get("uri") for pagina in doc for link in pagina.get_links()):
        return pdf
    for pagina in doc:
        palavras = pagina.get_text("words")
        k = 0
        while k < len(palavras):
            # o identificador pode quebrar em até 3 pedaços na coluna
            # estreita ("…-1-" / "000006/2026"): junta pedaços vizinhos
            for tamanho in (1, 2, 3):
                trecho = palavras[k:k + tamanho]
                if len(trecho) < tamanho:
                    break
                url = links.get("".join(w[4] for w in trecho))
                if url:
                    for w in trecho:
                        pagina.insert_link({"kind": pymupdf.LINK_URI,
                                            "uri": url,
                                            "from": pymupdf.Rect(w[:4])})
                    k += tamanho - 1
                    break
            k += 1
    return doc.tobytes(garbage=3, deflate=True)


def em_pdf(docx: bytes, branding: dict | None = None, *,
           tempo_limite: int = 120,
           links: dict[str, str] | None = None) -> bytes:
    """
    DOCX → PDF pelo conversor institucional (LibreOffice). Sem ele, ERRO
    explícito: degradar em silêncio para outro desenho é o que este
    relatório existe para não fazer.
    """
    from .. import export

    pdf = export._docx_em_pdf(docx, tempo_limite=tempo_limite)  # noqa: SLF001
    if not pdf:
        raise ErroDeRenderizacao(
            "A conversão para PDF não está disponível neste servidor agora "
            "(LibreOffice ausente ou tempo esgotado). Baixe o DOCX, que tem "
            "o mesmo conteúdo, ou tente de novo em instantes.")
    pdf = carimbar_links(pdf, links or {})
    return export._pdf_aplicar_marca(pdf, branding)  # noqa: SLF001


def relatorio_de_cotacao_pdf(view: CotacaoReportView) -> bytes:
    return em_pdf(relatorio_de_cotacao_docx(view),
                  dict(view.branding) if view.branding else None,
                  links=links_da_vista(view))


def relatorio_analitico_pdf(view: CotacaoReportView, *,
                            tempo_limite: int = 600) -> bytes:
    return em_pdf(relatorio_analitico_docx(view),
                  dict(view.branding) if view.branding else None,
                  tempo_limite=tempo_limite,
                  links=links_da_vista(view, analitico=True))
