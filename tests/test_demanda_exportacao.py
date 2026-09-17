"""
A planilha de juntada reproduz exatamente o que foi aprovado.

§16. Um export que arredonda, resume ou reordena vira um segundo
documento que diverge do processo — e, numa conferência posterior,
ninguém sabe qual dos dois é o oficial.
"""

from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
import pytest

from src.demanda import consolidacao, exportacao
from src.demanda.extracao import DFD, ItemDemanda


def _origem(sigla, itens, numero="1"):
    return consolidacao.Origem(secretaria=sigla, arquivo=f"{sigla}.pdf",
                               dfds=(DFD(numero=numero, itens=tuple(itens)),))


def _item(codigo, qtd, unidade="UNIDADE", descricao="CANETA"):
    return ItemDemanda(codigo=codigo, descricao=descricao,
                       quantidade=Decimal(str(qtd)), unidade=unidade)


@pytest.fixture
def consolidado():
    return consolidacao.consolidar([
        _origem("SEMAFI", [_item("111111", 30), _item("222222", 5, "CAIXA")]),
        _origem("SEMEC", [_item("111111", 70)]),
    ])


def _abrir(dados: bytes):
    return openpyxl.load_workbook(io.BytesIO(dados), data_only=True)


def test_a_planilha_tem_as_tres_abas(consolidado):
    wb = _abrir(exportacao.planilha_consolidada(consolidado))

    assert wb.sheetnames == ["Consolidado", "Ocorrências", "Resumo"]


def test_uma_coluna_por_secretaria_mais_o_total(consolidado):
    wb = _abrir(exportacao.planilha_consolidada(consolidado))
    cabecalho = [c.value for c in wb["Consolidado"][1]]

    assert cabecalho[:3] == ["Código do Item", "Descrição e Especificação", "Unidade"]
    assert "SEMAFI" in cabecalho and "SEMEC" in cabecalho
    assert cabecalho[-1] == "Total Geral"


def test_o_total_exportado_e_o_total_aprovado(consolidado):
    wb = _abrir(exportacao.planilha_consolidada(consolidado))
    aba = wb["Consolidado"]
    linha = next(r for r in aba.iter_rows(min_row=2, values_only=True)
                 if r[0] == "111111")

    assert linha[-1] == 100


def test_as_quantidades_sao_numero_e_nao_texto(consolidado):
    """
    Número como texto não soma, não ordena e não fecha conta — e a
    planilha existe para alguém conferir a soma.
    """
    wb = _abrir(exportacao.planilha_consolidada(consolidado))
    linha = next(r for r in wb["Consolidado"].iter_rows(min_row=2, values_only=True)
                 if r[0] == "111111")

    assert all(isinstance(c, (int, float)) for c in linha[3:] if c is not None)


def test_a_aba_de_ocorrencias_rastreia_cada_parcela(consolidado):
    wb = _abrir(exportacao.planilha_consolidada(consolidado))
    linhas = list(wb["Ocorrências"].iter_rows(min_row=2, values_only=True))

    assert len(linhas) == 3
    assert all(l[-1].endswith(".pdf") for l in linhas)
    assert all(l[-2] for l in linhas), "toda ocorrência aponta um DFD"


def test_a_soma_das_ocorrencias_reproduz_o_total(consolidado):
    """
    A prova do §12 dentro do próprio arquivo: quem receber o XLSX
    consegue refazer a conta sem o sistema.
    """
    wb = _abrir(exportacao.planilha_consolidada(consolidado))
    das_ocorrencias = sum(
        l[4] for l in wb["Ocorrências"].iter_rows(min_row=2, values_only=True)
        if l[1] == "111111")
    do_consolidado = next(
        r[-1] for r in wb["Consolidado"].iter_rows(min_row=2, values_only=True)
        if r[0] == "111111")

    assert das_ocorrencias == do_consolidado == 100


def test_o_resumo_declara_se_a_conta_fecha(consolidado):
    wb = _abrir(exportacao.planilha_consolidada(consolidado))
    resumo = {r[0]: r[1] for r in wb["Resumo"].iter_rows(min_row=2, values_only=True)}

    assert resumo["A soma fecha"] == "sim"
    assert resumo["Itens consolidados"] == 2


def test_o_resumo_nomeia_os_arquivos_nao_lidos():
    """
    O arquivo exportado é o que vai aos autos. Omitir ali que duas
    secretarias não foram lidas seria juntar ao processo um documento
    que afirma mais do que se sabe.
    """
    consolidado = consolidacao.consolidar([
        _origem("A", [_item("111111", 10)]),
        consolidacao.Origem(secretaria="SECULT", arquivo="secult.pdf",
                            status=consolidacao.IMAGEM),
    ])

    wb = _abrir(exportacao.planilha_consolidada(consolidado))
    valores = [r[1] for r in wb["Resumo"].iter_rows(min_row=2, values_only=True)]

    assert "secult.pdf" in valores


def test_o_resumo_nomeia_os_conflitos_pendentes():
    consolidado = consolidacao.consolidar([
        _origem("A", [_item("111111", 10, "CAIXA")]),
        _origem("B", [_item("111111", 10, "UNIDADE")]),
    ])

    wb = _abrir(exportacao.planilha_consolidada(consolidado))
    rotulos = [r[0] for r in wb["Resumo"].iter_rows(min_row=2, values_only=True)]

    assert any("111111" in str(r) for r in rotulos)
