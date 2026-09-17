import io
from docx import Document
from src import export


def test_export_docx_preserva_italico_negrito_listas_links_e_quebra():
    texto = '## Revisão\n\n**Ação** e *órgão*.\n\n1. Primeiro\n    - Interno\n        - Terceiro\n\n| Id | Dano |\n|---|---|\n| 1 | Atraso<br>Interrupção |\n\n[fonte](https://example.org)'
    doc = Document(io.BytesIO(export.gerar_docx("Teste",texto)))
    corpo = next(p for p in doc.paragraphs if "órgão" in p.text)
    assert any(r.text == "Ação" and r.bold for r in corpo.runs)
    assert any(r.text == "órgão" and r.italic for r in corpo.runs)
    assert next(p for p in doc.paragraphs if "Interno" in p.text).style.name == "GovDocs Item 2"
    assert next(p for p in doc.paragraphs if "Terceiro" in p.text).style.name == "GovDocs Item 3"
    assert doc.tables[0].cell(1,1).text == "Atraso\nInterrupção"
    assert any(r.target_ref == "https://example.org" for r in doc.part.rels.values())
