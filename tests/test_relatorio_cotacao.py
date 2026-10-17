"""
Relatório de Cotação, Relatório Analítico Integral e Memória Analítica —
o pacote de 16/10/2026 (`05_TESTES_ACEITE.md`).

O que estas provas guardam:

* os NÚMEROS não mudam: o que o documento imprime é o que o motor gravou
  (método legado continua mediana; média nova continua média);
* ausência não vira zero: item sem preço diz "Não estimado", e item não
  concluído não entra num total que pareça global;
* N/3 é a contagem do cálculo (`cesta.consideradas`), e referência
  excluída depois de selecionada não conta;
* as três saídas — compacto, analítico e XLSX — concordam em identificador,
  contagens e totais, porque leem a MESMA projeção;
* o PDF é legível: cabeçalho de tabela sem palavra partida, texto dentro
  da área útil, "Página X de Y", links reais e só os que existem nos dados.

Cenários sintéticos em `tests/relatorio_cotacao_apoio.py`; os preços
saem do motor real. Nada aqui chama rede nem IA.
"""

from __future__ import annotations

import copy
import io
import re
import zipfile
from dataclasses import FrozenInstanceError
from datetime import datetime, timezone
from decimal import Decimal

import pytest

from src.precos import cesta, relatorio
from src.precos import documento_cotacao as dc
from src.precos import relatorio_view as rv
from tests.relatorio_cotacao_apoio import cenario

MOMENTO = datetime(2026, 10, 16, 23, 30, tzinfo=timezone.utc)
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _vista(nome: str, **extras) -> rv.CotacaoReportView:
    return rv.projetar(*cenario(nome), emitido_em=MOMENTO, **extras)


def _texto_docx(conteudo: bytes) -> str:
    """Todo o texto do DOCX, um parágrafo por linha (inclui tabelas)."""
    from lxml import etree

    with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
        raiz = etree.fromstring(z.read("word/document.xml"))
    return "\n".join("".join(t.text or "" for t in p.iter(f"{_W}t"))
                     for p in raiz.iter(f"{_W}p"))


# ---------------------------------------------------------------------------
# Golden e invariantes do domínio (sem LibreOffice)
# ---------------------------------------------------------------------------
def test_golden_tres_referencias_media_e_total():
    v = _vista("minima")
    item = v.itens[0]
    assert item.metodo == "media"
    assert item.preco_unitario == Decimal("25.85")
    assert item.preco_total == Decimal("258.50")
    assert (item.consideradas, item.minimo) == (3, 3)
    assert v.subtotal_concluidos == Decimal("258.50") == v.valor_global
    texto = _texto_docx(dc.relatorio_de_cotacao_docx(v))
    assert "R$ 25,85" in texto and "R$ 258,50" in texto
    assert "Valor global estimado" in texto


def test_excluida_depois_de_selecionada_nao_conta():
    v = _vista("excluida_vira_2")
    item = v.itens[0]
    assert item.consideradas == 2 and not item.concluido
    assert v.subtotal_concluidos == 0 and v.valor_global is None
    texto = _texto_docx(dc.relatorio_de_cotacao_docx(v))
    assert "2/3" in texto and "Pendente" in texto
    assert "3/3" not in texto
    assert "Valor global estimado" not in texto
    assert "valor global não consolidado" in texto


@pytest.mark.parametrize("nome", ["sem_preco", "fontes_ausentes"])
def test_sem_preco_nao_vira_zero(nome):
    v = _vista(nome)
    assert all(i.preco_unitario is None for i in v.itens)
    texto = _texto_docx(dc.relatorio_de_cotacao_docx(v))
    assert "R$ 0,00" not in texto
    assert "Não estimado" in texto and "0/3" in texto


def test_valor_provisorio_nao_entra_no_quadro_nem_no_subtotal():
    """Incompleto COM valor calculado: aparece como provisório, não soma."""
    v = _vista("insuficiente")
    item = v.itens[0]
    assert item.preco_unitario is not None and item.valor_provisorio
    texto = _texto_docx(dc.relatorio_de_cotacao_docx(v))
    assert f"{dc.moeda(item.preco_unitario)} (provisório)" in texto
    assert "Não consolidado" in texto
    assert v.subtotal_concluidos == 0


def test_subtotal_e_soma_exata_dos_concluidos():
    p, itens, refs = cenario("processo_210")
    v = rv.projetar(p, itens, refs, emitido_em=MOMENTO)
    esperado = sum(Decimal(i["preco_total"]) for i in itens
                   if i["estado"] == "complete")
    assert v.subtotal_concluidos == esperado
    assert v.concluidos + v.nao_concluidos == 210
    assert v.valor_global is None          # há pendentes


def test_metodo_legado_continua_mediana():
    p, itens, refs = cenario("historico_5_871")
    v = rv.projetar(p, itens, refs, emitido_em=MOMENTO)
    assert {i.metodo for i in v.itens} == {"mediana"}
    for item, bruto in zip(v.itens, itens):
        assert item.preco_unitario == Decimal(bruto["preco_estimado"])
    assert "regra anterior a 15/10/2026" in v.regra_rotulo
    texto = _texto_docx(dc.relatorio_de_cotacao_docx(v))
    assert "mediana" in texto and "média aritmética simples como" not in texto


def test_contagem_e_a_do_calculo():
    p, itens, refs = cenario("historico_5_871")
    v = rv.projetar(p, itens, refs, emitido_em=MOMENTO)
    for item in v.itens:
        assert item.consideradas == len(cesta.consideradas(refs[item.id]))
        assert len(item.considerada_lista) == item.consideradas
    assert v.total_consideradas == 871


def test_marcada_sem_preco_normalizado_fica_fora_e_sem_link_fabricado():
    p, itens, refs = cenario("minima")
    linha = copy.deepcopy(refs["item-001"][0])
    linha.update({"id": "ref-sem-unidade", "valor_unitario_normalizado": None,
                  "bruto": {"id": "x"}, "fonte_id": "pncp"})
    refs["item-001"].append(linha)
    v = rv.projetar(p, itens, refs, emitido_em=MOMENTO)
    fora = [r for r in v.itens[0].referencias if r.id == "ref-sem-unidade"][0]
    assert not fora.considerada
    assert fora.situacao == "Marcada, fora do cálculo"
    assert fora.link is None
    assert v.itens[0].consideradas == 3


def test_volume_exibe_amostra_declarada_e_anexo_tem_todas():
    for nome, n, excluidas in (("grande_196", 196, 20),
                               ("grande_317", 317, 30)):
        v = _vista(nome)
        item = v.itens[0]
        assert item.consideradas == n and len(item.exibidas) == 5
        compacto = _texto_docx(dc.relatorio_de_cotacao_docx(v))
        assert f"exibidas 5 de {n}" in compacto
        assert f"Todas as {n} referências consideradas compõem o cálculo" \
            in compacto
        analitico = _texto_docx(dc.relatorio_analitico_docx(v))
        for r in item.referencias:
            assert r.id_externo in analitico
        assert item.total_referencias == n + excluidas


def test_representativas_cobrem_a_dispersao_sem_mudar_a_cesta():
    v = _vista("grande_196")
    item = v.itens[0]
    valores = sorted(r.valor for r in item.considerada_lista)
    exibidas = [r.valor for r in item.exibidas]
    assert exibidas[0] == valores[0] and exibidas[-1] == valores[-1]
    assert exibidas == sorted(exibidas)
    assert item.consideradas == 196          # a amostra é só de exibição


def test_tres_saidas_concordam():
    """Compacto, analítico e XLSX: identificador, N por item e totais."""
    from openpyxl import load_workbook

    p, itens, refs = cenario("pendencias")
    v = rv.projetar(p, itens, refs, emitido_em=MOMENTO)
    compacto = _texto_docx(dc.relatorio_de_cotacao_docx(v))
    analitico = _texto_docx(dc.relatorio_analitico_docx(v))
    planilha = load_workbook(io.BytesIO(
        relatorio.xlsx_analitico(p, itens, refs)))
    ficha = {linha[0].value: linha[1].value
             for linha in planilha["Identificação"].iter_rows(min_row=2)}
    assert v.hash_decisorio in compacto and v.hash_decisorio in analitico
    assert ficha["Identificador da versão"] == v.hash_decisorio
    aba = planilha["Itens"]
    cabecalho = [c.value for c in aba[1]]
    col_n = cabecalho.index("Consideradas no cálculo")
    col_total = cabecalho.index("Valor total")
    por_numero = {linha[0].value: linha for linha in aba.iter_rows(min_row=2)}
    for item in v.itens:
        linha = por_numero[item.numero]
        assert linha[col_n].value == item.consideradas
        assert f"{item.consideradas}/{item.minimo}" in compacto
        assert f"{item.consideradas}/{item.minimo}" in analitico
        if item.preco_total is not None:
            assert Decimal(str(linha[col_total].value)) == item.preco_total
    referencias = planilha["Referências"]
    assert referencias.max_row - 1 == v.total_referencias
    cab = [c.value for c in referencias[1]]
    links = {linha[cab.index("Link de verificação")].value
             for linha in referencias.iter_rows(min_row=2)} - {None, ""}
    ids = {linha[cab.index("ID da referência")].value
           for linha in referencias.iter_rows(min_row=2)}
    assert links == {r.link for i in v.itens for r in i.referencias if r.link}
    assert ids == {r.id for i in v.itens for r in i.referencias}


def test_perfil_real_sem_conformidade_automatica_a_norma_federal():
    texto = _texto_docx(dc.relatorio_de_cotacao_docx(_vista("minima")))
    assert "Lei nº 14.133/2021" in texto
    assert "Em conformidade com a Instrução Normativa" not in texto


@pytest.mark.parametrize("proibido", [
    "Banco de Preços", "bancodeprecos", "Certificado", "Código Validação",
    "CertificadoAutenticado", "QR Code", "qrcode", "05.193.057",
])
def test_nada_do_documento_comercial_de_referencia(proibido):
    for gerar in (dc.relatorio_de_cotacao_docx, dc.relatorio_analitico_docx):
        assert proibido not in _texto_docx(gerar(_vista("pendencias")))


def test_projecao_nao_altera_os_dados_e_e_imutavel():
    p, itens, refs = cenario("pendencias")
    antes = copy.deepcopy((p, itens, refs))
    v = rv.projetar(p, itens, refs, emitido_em=MOMENTO)
    assert (p, itens, refs) == antes
    with pytest.raises(FrozenInstanceError):
        v.nome = "outro"           # type: ignore[misc]
    with pytest.raises(TypeError):
        v.itens[0].estatisticas["media"] = Decimal("1")  # type: ignore[index]


def test_identificador_nao_depende_da_emissao():
    a = rv.projetar(*cenario("minima"), emitido_em=MOMENTO)
    b = rv.projetar(*cenario("minima"),
                    emitido_em=datetime(2027, 1, 1, tzinfo=timezone.utc))
    assert a.hash_decisorio == b.hash_decisorio


def test_cotacao_sem_itens_e_documento_explicito():
    texto = _texto_docx(dc.relatorio_de_cotacao_docx(_vista("vazio")))
    assert "não tem itens" in texto


def test_especificacao_extensa_aparece_inteira():
    p, itens, refs = cenario("extensa")
    assert len(itens[0]["descricao"]) > 500
    texto = _texto_docx(dc.relatorio_de_cotacao_docx(
        rv.projetar(p, itens, refs, emitido_em=MOMENTO)))
    assert " ".join(itens[0]["descricao"].split()) in texto.replace("\n", " ")


def test_excecao_formal_aparece_com_fundamento_e_aprovacao():
    texto = _texto_docx(dc.relatorio_de_cotacao_docx(_vista("pendencias")))
    assert "Exceção à regra dos três" in texto
    assert "mercado restrito" in texto
    assert "Autoridade competente (fictícia)" in texto


def test_timbrado_e_do_tenant_informado():
    import base64

    from PIL import Image

    imagem = io.BytesIO()
    Image.new("RGB", (800, 120), (20, 60, 120)).save(imagem, format="PNG")
    marca_a = {"orgao": "PREFEITURA A (fictícia)",
               "cabecalho_img": base64.b64encode(imagem.getvalue()).decode()}
    docx_a = dc.relatorio_de_cotacao_docx(_vista("minima", branding=marca_a))
    docx_b = dc.relatorio_de_cotacao_docx(_vista(
        "minima", branding={"orgao": "PREFEITURA B (fictícia)"}))
    assert "PREFEITURA A (fictícia)" in _texto_docx(docx_a)
    assert "PREFEITURA A" not in _texto_docx(docx_b)
    with zipfile.ZipFile(io.BytesIO(docx_a)) as z:
        assert any(n.startswith("word/media/") for n in z.namelist())
    with zipfile.ZipFile(io.BytesIO(docx_b)) as z:
        assert not any(n.startswith("word/media/") for n in z.namelist())
        cabecalhos = " ".join(z.read(n).decode("utf-8") for n in z.namelist()
                              if n.startswith("word/header"))
        assert "PREFEITURA B (FICTÍCIA)" in cabecalhos


@pytest.mark.parametrize("descricao, esperado", [
    ("PAPEL A4 BRANCO. ESPECIFICAÇÃO: papel sulfite", "PAPEL A4 BRANCO"),
    ("GRAMPO 24/10 CX COM 5000 ESPECIFICAÇÃO: aço", "GRAMPO 24/10 CX COM 5000"),
    ("METRÔNOMO DIGITAL", "METRÔNOMO DIGITAL"),
    ("", rv.NAO_INFORMADO),
])
def test_titulo_curto(descricao, esperado):
    assert rv.titulo_curto(descricao) == esperado


def test_titulo_curto_longo_corta_em_palavra():
    titulo = rv.titulo_curto("CADEIRA " * 30)
    assert len(titulo) <= 64 and titulo.endswith("…")
    assert not titulo[:-1].endswith(" ")


def test_markdown_legado_inalterado_pela_extracao_de_funcoes():
    """`completo`/`resumido` continuam formatando a trilha e os alertas."""
    p, itens, refs = cenario("pendencias")
    texto = relatorio.completo(p, itens, refs, emitido_em=MOMENTO)
    assert "| Fonte | Enquadramento | Referências | Consideradas |" in texto
    assert "- **Itens sem preço formado (2):** 2, 3." in texto
    assert "Itens ainda em revisão (1)" in texto


# ---------------------------------------------------------------------------
# PDF real (LibreOffice) — geometria, paginação, links
# ---------------------------------------------------------------------------
def _spans(pdf):
    for numero, pagina in enumerate(pdf, start=1):
        for bloco in pagina.get_text("dict")["blocks"]:
            for linha in bloco.get("lines", []):
                for span in linha["spans"]:
                    if span["text"].strip():
                        yield numero, pagina, span


def _conferir_geometria(pdf, *, margem_pt: float) -> None:
    for numero, pagina, span in _spans(pdf):
        x0, _, x1, _ = span["bbox"]
        assert x0 >= margem_pt - 3, (numero, span["text"], span["bbox"])
        assert x1 <= pagina.rect.width - margem_pt + 3, (
            numero, span["text"], span["bbox"])


def _conferir_paginacao(pdf) -> None:
    total = len(pdf)
    for numero, pagina in enumerate(pdf, start=1):
        texto = " ".join(pagina.get_text("text").split())
        assert f"Página {numero} de {total}" in texto, numero


def _cabecalhos_inteiros(pdf, titulos: list[str]) -> None:
    """Cada título de coluna aparece INTEIRO numa linha (sem Méto/do)."""
    linhas = set()
    for pagina in pdf:
        for bloco in pagina.get_text("dict")["blocks"]:
            for linha in bloco.get("lines", []):
                linhas.add("".join(s["text"] for s in linha["spans"]).strip())
    for titulo in titulos:
        assert titulo in linhas, titulo


@pytest.fixture(scope="module")
def pdf_historico():
    import pymupdf

    from tests.conftest import exigir_motor_institucional

    exigir_motor_institucional()
    v = _vista("historico_5_871")
    return v, pymupdf.open(stream=dc.relatorio_de_cotacao_pdf(v))


def test_pdf_compacto_5_itens_871_referencias(pdf_historico):
    v, pdf = pdf_historico
    assert len(pdf) <= 15
    texto = " ".join(p.get_text("text") for p in pdf)
    assert "RELATÓRIO DE COTAÇÃO DE PREÇOS" in texto
    _conferir_geometria(pdf, margem_pt=1.6 / 2.54 * 72)
    _conferir_paginacao(pdf)
    _cabecalhos_inteiros(pdf, ["Nº", "Item", "Qtd. / Unid.", "Válidas",
                               "Unitário", "Total", "Situação",
                               "Fonte / identificador", "Preço unit."])
    for proibido in ("Méto\ndo", "Situaç\não", "Unitá\nrio"):
        assert proibido not in texto


def test_pdf_links_sao_reais_e_so_os_dos_dados(pdf_historico):
    v, pdf = pdf_historico
    nos_dados = {r.link for i in v.itens for r in i.referencias if r.link}
    no_pdf = {l["uri"] for p in pdf for l in p.get_links() if l.get("uri")}
    assert no_pdf and no_pdf <= nos_dados


def test_pdf_com_pendencias_nao_consolida_global(motor_institucional):
    import pymupdf

    v = _vista("pendencias")
    pdf = pymupdf.open(stream=dc.relatorio_de_cotacao_pdf(v))
    texto = " ".join(" ".join(p.get_text("text").split()) for p in pdf)
    assert "Subtotal dos itens concluídos (2 de 5)" in texto
    assert "valor global não consolidado" in texto
    assert "Valor global estimado" not in texto
    assert "R$ 0,00" not in texto
    _conferir_geometria(pdf, margem_pt=1.6 / 2.54 * 72)
    _conferir_paginacao(pdf)


def test_pdf_analitico_paisagem_com_todas_as_referencias(motor_institucional):
    import pymupdf

    v = _vista("grande_196")
    pdf = pymupdf.open(stream=dc.relatorio_analitico_pdf(v))
    assert pdf[0].rect.width > pdf[0].rect.height
    texto = " ".join(p.get_text("text") for p in pdf)
    for r in v.itens[0].referencias:
        assert r.id_externo in texto
    _conferir_paginacao(pdf)
    _conferir_geometria(pdf, margem_pt=1.6 / 2.54 * 72)
    # cabeçalho da tabela repetido em toda página com tabela
    for pagina in list(pdf)[1:-1]:
        assert "Fonte · inciso · natureza" in pagina.get_text("text")
    _cabecalhos_inteiros(pdf, ["Situação", "Compat.", "Preço unit.",
                               "Identificador", "Motivo"])
    assert re.search(r"Página \d+ de \d+", texto)


def test_objeto_e_criterios_de_busca_estao_no_relatorio():
    """Os 22 pontos do relatório anterior: §2 e §8 não podem sumir."""
    texto = _texto_docx(dc.relatorio_de_cotacao_docx(_vista("minima")))
    assert "Objeto: Aquisição de materiais de expediente" in texto
    assert "últimos 180 dias contados da data-base" in texto
    assert "recorte geográfico: PA" in texto


def test_pdf_bloco_do_item_nao_se_parte(motor_institucional):
    """Faixa, rótulos e valores dos indicadores na MESMA página."""
    import pymupdf

    v = _vista("processo_210")
    pdf = pymupdf.open(stream=dc.relatorio_de_cotacao_pdf(v))
    faixa, rotulo, valor = {}, [], []
    for numero, pagina, span in _spans(pdf):
        texto = span["text"].strip()
        achado = re.match(r"ITEM (\d{2,3}) ·", texto)
        if achado and span["size"] >= 9.5:
            faixa.setdefault(int(achado.group(1)), numero)
        if texto == "Valor do item":
            rotulo.append(numero)
        if re.fullmatch(r"\d+/\d+", texto) and span["size"] >= 9.5:
            valor.append(numero)
    assert len(faixa) == len(v.itens) == len(rotulo) == len(valor)
    for item, pagina_rotulo, pagina_valor in zip(v.itens, rotulo, valor):
        assert faixa[item.numero] == pagina_rotulo == pagina_valor, item.numero


def test_links_carimbados_quando_o_conversor_os_descarta(pdf_historico):
    """
    LibreOffice 7.4 (produção) descarta os hyperlinks do DOCX. Sem link
    nenhum no PDF, `carimbar_links` grava cada um sobre o identificador
    impresso — só com URL que a projeção tem.
    """
    import pymupdf

    v, pdf = pdf_historico
    sem_links = pymupdf.open(stream=pdf.tobytes())
    for pagina in sem_links:
        for link in pagina.get_links():
            pagina.delete_link(link)
    cru = sem_links.tobytes()
    assert not any(p.get_links() for p in pymupdf.open(stream=cru))
    carimbado = pymupdf.open(stream=dc.carimbar_links(cru, dc.links_da_vista(v)))
    uris = [l["uri"] for p in carimbado for l in p.get_links()]
    assert uris and set(uris) <= set(dc.links_da_vista(v).values())
    # cada link cai sobre o identificador cuja URL ele carrega
    por_id = dc.links_da_vista(v)
    for pagina in carimbado:
        for link in pagina.get_links():
            palavra = pagina.get_textbox(link["from"]).strip()
            assert por_id.get(palavra) == link["uri"]
    # com links já presentes, não duplica
    com_links = pdf.tobytes()
    assert dc.carimbar_links(com_links, por_id) is com_links
