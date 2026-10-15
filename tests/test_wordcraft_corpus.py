import importlib.util
import hashlib
import io
import json
from pathlib import Path
import shutil
import zipfile

import pytest

SPEC = importlib.util.spec_from_file_location(
    "wordcraft_corpus", Path(__file__).resolve().parents[1] / "scripts/wordcraft_corpus.py")
corpus = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(corpus)


def test_210_itens_totais_e_fontes_sinteticas():
    table, oracle = corpus.tabela_sintetica()
    assert oracle["items"] == len(set(oracle["codes"])) == 210
    assert all(table.count(code) == 1 for code in oracle["codes"])
    assert table.count("https://example.invalid/precos/") == 210
    assert len(oracle["line_totals_decimal"]) == 210


def test_corpus_docx_preserva_tabela_links_identidade_assinaturas(tmp_path):
    manifest = corpus.generate(tmp_path)
    assert manifest["synthetic_only"]
    cases = manifest["cases"]
    assert cases[0]["hyperlinks"] == 1
    assert cases[1]["table_rows"] == [212]  # header + 210 items + global total
    assert cases[1]["hyperlinks"] == 210
    assert cases[2]["headers"] == cases[2]["footers"] == 1
    inspected = corpus.inspect_docx((tmp_path / "03-timbrado-assinaturas.docx").read_bytes())
    assert "Pessoa Sintética A" in inspected["text"]
    assert "Pessoa Sintética B" in inspected["text"]


def test_pdf_nao_aceita_fallback_como_referencia(tmp_path, monkeypatch):
    monkeypatch.setattr(corpus, "motor_pdf", lambda: "fpdf2")
    with pytest.raises(ValueError, match="LibreOffice"):
        corpus.generate(tmp_path / "corpus", pdf=True)
    assert not (tmp_path / "corpus").exists()


def test_fixtures_versionadas_sao_sinteticas_e_byte_identicas_ao_manifesto():
    inputs = Path(__file__).resolve().parent / "fixtures/wordcraft"
    manifest = corpus.verify_inputs(inputs)
    assert manifest["source_run"] == 38620142140
    assert manifest["oracle"]["items"] == 210
    table = corpus.text_blocks((inputs / "02-tabela-210-itens.docx").read_bytes())
    assert len(table["links"]) == 210
    assert all(code in "\n".join(table["body"]) for code in manifest["oracle"]["codes"])


def test_comparador_recusa_texto_perdido_no_retorno(tmp_path):
    inputs = Path(__file__).resolve().parent / "fixtures/wordcraft"
    manifest = corpus.verify_inputs(inputs)
    for item in manifest["files"]:
        target = tmp_path / Path(item["path"]).stem
        target.mkdir()
        (target / "roundtrip.docx").write_bytes((inputs / item["path"]).read_bytes())
        (target / "model.txt").write_text("\n".join(manifest["oracle"]["codes"]), encoding="utf-8")
    assert len(corpus.compare_roundtrip(inputs, tmp_path)["cases"]) == 3
    broken = corpus.gerar_docx("Ensaio", "Texto deliberadamente perdido")
    (tmp_path / "01-texto-listas-links/roundtrip.docx").write_bytes(broken)
    with pytest.raises(ValueError, match="changed semantic"):
        corpus.compare_roundtrip(inputs, tmp_path)


@pytest.fixture(scope="module")
def full_generated(tmp_path_factory):
    output = tmp_path_factory.mktemp("g3-full-corpus")
    manifest = corpus.generate_full(output)
    return output, manifest


def full_variant(full_generated, case_id, number=0):
    output, manifest = full_generated
    case = next(item for item in manifest["cases"] if item["id"] == case_id)
    variant = case["variants"][number]
    oracle = json.loads((output / variant["oracle"]).read_text(encoding="utf-8"))
    return case, variant, oracle, (output / variant["docx"]).read_bytes()


def test_full_corpus_has_twelve_seeds_and_keeps_acceptance_pending(full_generated):
    output, manifest = full_generated
    assert manifest["schema"] == "gc-wordcraft-full-corpus/v1"
    assert {case["id"] for case in manifest["cases"]} == {f"C{i:02d}" for i in range(1, 13)}
    assert sum(len(case["variants"]) for case in manifest["cases"]) == 16
    assert corpus.verify_full_inputs(output) == manifest
    for case in manifest["cases"]:
        assert set(case["workflow_acceptance"].values()) <= {"NOT_TESTED"}
        for variant in case["variants"]:
            assert variant["fidelity"] == {"wordcraft_roundtrip": "NOT_RUN",
                "visual_review": "PENDING", "libreoffice_pdf": "NOT_REQUESTED"}
    # The separate generation never changes the reviewed three-file inventory.
    assert len(corpus.verify_inputs(Path(__file__).parent / "fixtures/wordcraft")["files"]) == 3


def test_dfd_uses_five_items_decimal_oracle_and_synthetic_eligible_snapshots(full_generated):
    _, variant, oracle, data = full_variant(full_generated, "C01")
    assert oracle["items"] == 5
    assert oracle["line_totals_decimal"] == ["0.10", "0.40", "0.90", "5002.20", "4999.95"]
    assert oracle["global_total_decimal"] == "10003.55"
    assert len(oracle["clause_titles"]) == 9
    assert len(oracle["signatarios_snapshot"]) == len(oracle["signature_names"]) == 2
    assert oracle["portaria_snapshot"]["numero"] == "QA-001"
    assert oracle["institutional_origin"] == "secretaria"
    text = corpus.inspect_docx(data)["text"]
    assert all(name in text for name in oracle["signature_names"])
    assert "Portaria nº QA-001/2026" in text
    assert variant["origin"] == "canonical_product_export"


def test_etp_long_has_all_profile_headings_and_unique_ordered_paragraphs(full_generated):
    _, _, oracle, data = full_variant(full_generated, "C02")
    from src import perfis
    assert oracle["clause_titles"] == [
        f"{clause['n']}. {clause['titulo']}" for clause in perfis.clausulas_aplicaveis("etp", srp=True)]
    assert len(oracle["clause_titles"]) == 18
    assert len(oracle["markers"]) == 90
    text = corpus.inspect_docx(data)["text"]
    assert len(text.split()) > 6000
    assert [text.index(marker) for marker in oracle["markers"]] == sorted(
        text.index(marker) for marker in oracle["markers"])


def test_risk_fixture_has_real_horizontal_and_vertical_merges(full_generated):
    _, variant, oracle, data = full_variant(full_generated, "C03")
    structure = corpus.inspect_full_docx(data)
    assert structure["horizontal_spans"] == [6]
    assert structure["vertical_merge_restarts"] == 3
    assert structure["vertical_merge_continuations"] == 9
    assert structure["page_break_before"] >= oracle["explicit_page_break_before"]
    assert len(oracle["markers"]) == 12
    assert variant["origin"] == "canonical_export_plus_python_docx_QA"


def test_tr_thirty_items_has_independent_known_total_and_checks_individual_cells(full_generated):
    _, _, oracle, data = full_variant(full_generated, "C04")
    assert oracle["items"] == 30
    assert oracle["global_total_decimal"] == "1085.85"
    assert oracle["global_total_display"] == "R$ 1.085,85"
    assert oracle["line_totals_decimal"][0] == "20.36"
    assert oracle["line_totals_decimal"][-1] == "13.95"
    from docx import Document
    document = Document(io.BytesIO(data))
    document.tables[0].cell(1, 5).text = "R$ 999.999,99"
    altered = io.BytesIO()
    document.save(altered)
    with pytest.raises(ValueError, match="monetary cells"):
        corpus._verificar_entrada_docx(altered.getvalue(), oracle)


def test_bad_existing_calculator_cannot_define_the_expected_oracle(monkeypatch):
    calculate = corpus.planilha.calcular

    def wrong(source):
        rows, total = calculate(source)
        rows[0]["valor_total"] += 1
        return rows, total

    monkeypatch.setattr(corpus.planilha, "calcular", wrong)
    source = [{"codigo": "QA-ONE", "descricao": "Entrada fictícia", "unidade": "UN",
               "quantidade": "3", "valor_unitario": "0,10",
               "fonte": "https://example.invalid/one"}]
    with pytest.raises(ValueError, match="independent full-corpus oracle"):
        corpus.tabela_com_oraculo(source)


def test_edital_has_real_bookmarks_internal_remissions_and_external_links(full_generated):
    _, _, oracle, data = full_variant(full_generated, "C05")
    inspection = corpus.inspect_full_docx(data)
    assert inspection["bookmarks"] == inspection["internal_links"] == oracle["bookmarks"]
    assert len(oracle["bookmarks"]) == len(oracle["clause_titles"]) == 32
    assert len(oracle["external_links"]) == 8
    assert corpus.text_blocks(data)["links"] == sorted(oracle["external_links"])
    assert "Índice de teste para 32." in inspection["text"]


def test_arp_positive_and_negative_use_canonical_process_condition(full_generated):
    _, _, positive, included = full_variant(full_generated, "C06", 0)
    case, _, negative, excluded = full_variant(full_generated, "C06", 1)
    assert positive["conditional_arp_expected"] is True
    assert negative["conditional_arp_expected"] is False
    assert "C06-ARP-ONLY" in corpus.inspect_docx(included)["text"]
    assert "C06-ARP-ONLY" not in corpus.inspect_docx(excluded)["text"]
    assert "C06-EDITAL-ONLY" in corpus.inspect_docx(excluded)["text"]
    assert case["workflow_acceptance"] == {"OPERATIONAL_CONDITIONAL_EMISSION_NOT_TESTED": "NOT_TESTED"}


def test_branding_has_original_logo_snapshot_and_municipal_fallback(full_generated):
    _, _, secretary, data = full_variant(full_generated, "C07", 0)
    _, _, municipality, fallback = full_variant(full_generated, "C07", 1)
    assert secretary["institutional_origin"] == "secretaria"
    assert municipality["institutional_origin"] == "municipio"
    assert secretary["logo_sha256"] in corpus.inspect_full_docx(data)["media_sha256"]
    assert secretary["logo_sha256"] == hashlib.sha256(corpus._logo_sintetico()).hexdigest()
    assert "MUNICÍPIO SINTÉTICO A" in "\n".join(corpus.text_blocks(fallback)["headers"])
    institution = corpus.instituicao_sintetica()
    frozen = dict(institution["branding_snapshot"])
    institution["secretary"]["cabecalho"] = "CADASTRO POSTERIOR DO ENSAIO"
    assert institution["branding_snapshot"] == frozen
    assert institution["approval_status"] == "FIXTURE_ONLY_NOT_DATABASE_APPROVAL"


def test_full_scale_keeps_all_210_codes_and_the_original_independent_total(full_generated):
    _, _, oracle, data = full_variant(full_generated, "C08")
    assert oracle["items"] == 210
    assert oracle["global_total_decimal"] == "92900.30"
    text = corpus.inspect_docx(data)["text"]
    assert all(text.count(code) == 1 for code in oracle["codes"])


def test_opinion_findings_and_governance_cases_never_report_simulated_acceptance(full_generated):
    output, _ = full_generated
    case, _, opinion, _ = full_variant(full_generated, "C09")
    from src.blocos import hash_texto
    assert [finding["nature"] for finding in opinion["findings"]] == ["objetiva", "ambigua", "conflitante"]
    assert opinion["source_markdown_hash"] == hash_texto((output / "C09/alvo-sintetico.md").read_text(encoding="utf-8"))
    assert set(case["workflow_acceptance"].values()) == {"NOT_TESTED"}
    assert opinion["status"] == "EXPECTED_ONLY_NOT_PRODUCT_ACCEPTANCE"
    historic, original, _, original_bytes = full_variant(full_generated, "C11", 0)
    _, revised, _, revised_bytes = full_variant(full_generated, "C11", 1)
    assert original["sha256_docx"] != revised["sha256_docx"]
    assert original_bytes != revised_bytes
    assert historic["workflow_acceptance"] == {"APPROVAL_AND_IMMUTABILITY_NOT_TESTED": "NOT_TESTED"}
    tenants, _, tenant_a, bytes_a = full_variant(full_generated, "C12", 0)
    _, _, tenant_b, bytes_b = full_variant(full_generated, "C12", 1)
    assert tenant_a["tenant_id"] != tenant_b["tenant_id"]
    assert tenant_a["secretaria_id"] != tenant_b["secretaria_id"]
    assert "QA-TENANT-B-ONLY" not in corpus.inspect_docx(bytes_a)["text"]
    assert "QA-TENANT-A-ONLY" not in corpus.inspect_docx(bytes_b)["text"]
    assert set(tenants["workflow_acceptance"].values()) == {"NOT_TESTED"}


def test_external_model_declares_and_contains_nested_table_merge_and_page_field(full_generated):
    _, variant, oracle, data = full_variant(full_generated, "C10")
    inspection = corpus.inspect_full_docx(data)
    assert inspection["table_rows"] == [2, 2]
    assert inspection["horizontal_spans"] == [2]
    assert inspection["explicit_page_breaks"] == 1
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert b'PAGE' in archive.read("word/footer1.xml")
    assert variant["origin"] == "canonical_export_plus_python_docx_QA"
    assert "nested_table" in oracle["declared_features"]


def test_full_generator_never_queries_database_or_resolves_remote_sources(monkeypatch, tmp_path):
    from src import db

    def forbidden(*args, **kwargs):
        pytest.fail("full corpus must use only local pure fixture data")

    monkeypatch.setattr(db, "_cliente", forbidden)
    monkeypatch.setattr(db, "cliente_do_usuario", forbidden)
    monkeypatch.setattr(db, "listar_secretarias", forbidden)
    import socket
    monkeypatch.setattr(socket, "create_connection", forbidden)
    assert len(corpus.generate_full(tmp_path / "offline")["cases"]) == 12


def test_full_pdf_refuses_fallback_and_failed_lo_before_writing(tmp_path, monkeypatch):
    output = tmp_path / "missing-lo"
    monkeypatch.setattr(corpus, "motor_pdf", lambda: "fpdf2")
    with pytest.raises(ValueError, match="real institutional LibreOffice"):
        corpus.generate_full(output, pdf=True)
    assert not output.exists()
    from src import export
    monkeypatch.setattr(corpus, "motor_pdf", lambda: "libreoffice")
    monkeypatch.setattr(export, "_docx_em_pdf", lambda content: None)
    with pytest.raises(ValueError, match="no PDF fallback"):
        corpus.generate_full(output, pdf=True)
    assert not output.exists()


def test_full_generation_refuses_the_immutable_g1_directory_and_existing_content(tmp_path):
    with pytest.raises(ValueError, match="immutable G1"):
        corpus.generate_full(Path(__file__).parent / "fixtures/wordcraft")
    occupied = tmp_path / "occupied"
    occupied.mkdir()
    sentinel = occupied / "original.txt"
    sentinel.write_bytes(b"keep the original")
    with pytest.raises(ValueError, match="fresh full-corpus"):
        corpus.generate_full(occupied)
    assert sentinel.read_bytes() == b"keep the original"


def test_full_input_digest_check_rejects_modified_docx(full_generated, tmp_path):
    output, manifest = full_generated
    target = tmp_path / "tampered"
    shutil.copytree(output, target)
    first = target / manifest["cases"][0]["variants"][0]["docx"]
    first.write_bytes(first.read_bytes() + b"tampered")
    with pytest.raises(ValueError, match="digest differs"):
        corpus.verify_full_inputs(target)


@pytest.mark.parametrize("column,value", [(3, "99999"), (5, "R$ 999.999,99")])
def test_full_scale_oracle_rejects_altered_quantity_and_global_total(full_generated, column, value):
    _, _, oracle, data = full_variant(full_generated, "C08")
    from docx import Document
    document = Document(io.BytesIO(data))
    row = 1 if column == 3 else len(document.tables[0].rows) - 1
    document.tables[0].cell(row, column).text = value
    altered = io.BytesIO()
    document.save(altered)
    with pytest.raises(ValueError, match="quantity differs|global total differs"):
        corpus._verificar_entrada_docx(altered.getvalue(), oracle)
