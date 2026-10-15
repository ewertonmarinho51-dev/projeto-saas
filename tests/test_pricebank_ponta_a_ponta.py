"""
Banco de Preços de ponta a ponta no serviço (Gate 2, 15/10/2026).

Consulta manual → pesquisa salva → exclusão com motivo → preço refeito →
item confirmado → relatório completo, quadro, XLSX e PDF → aplicação ao
processo. Tudo pelo repositório REAL sobre o banco falso dos testes de
preços: a consulta manual não tem caminho próprio de relatório nem de
aplicação, usa os da pesquisa automática.
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from src import db, export
from src.precos import aplicacao, busca, cesta, perfil, relatorio
from src.precos import repositorio as repo
from src.precos.estados import EstadoItem
from tests.conftest import ClientePrecosFalso
from src.precos.modelo import NaturezaValor
from tests.test_precos_fase4 import FONTE_OFICIAL, _FonteFalsa, _referencia

ITEM = {"descricao": "CANETA ESFEROGRAFICA AZUL", "unidade": "UNIDADE",
        "quantidade": "100"}


def _homologada(ident, valor, natureza=NaturezaValor.HOMOLOGADO):
    ref = _referencia(FONTE_OFICIAL, ident, valor)
    ref.natureza_valor = natureza
    return ref


def _fonte(*extras):
    valores = ("20.00", "22.00", "24.00", "200.00")
    return _FonteFalsa(FONTE_OFICIAL, [
        _homologada(f"r{i}", v) for i, v in enumerate(valores)] + list(extras))


def _salvar(monkeypatch):
    cliente = ClientePrecosFalso()
    monkeypatch.setattr(db, "cliente_do_usuario", lambda: cliente)
    coleta = busca.coletar(dict(ITEM, id="consulta", numero=1), [_fonte()])
    chaves = {r.chave_dedupe for r, _ in coleta.referencias}
    salvo = cesta.salvar_consulta_como_pesquisa(
        repositorio=repo, nome="Banco de Preços — canetas", item=ITEM,
        coleta=coleta, selecionados=chaves, perfil=perfil.LEI_14133,
        auth_user_id="u1", idempotency_key="banco:t1")
    return cliente, coleta, salvo


def test_consulta_manual_ate_o_arquivo_exportado_e_o_processo(monkeypatch):
    cliente, _, salvo = _salvar(monkeypatch)
    pesquisa = repo.obter_pesquisa(salvo["pesquisa_id"])
    [item] = repo.listar_itens(salvo["pesquisa_id"])
    assert Decimal(str(item["preco_estimado"])) == Decimal("66.50")

    # exclusão do discrepante, com motivo, refaz o preço
    linhas = repo.listar_referencias(str(item["id"]))
    discrepante = next(r for r in linhas
                       if Decimal(str(r["valor_unitario_normalizado"])) == 200)
    repo.reclassificar_referencia(str(discrepante["id"]), "rejected",
                                  "preço 9× acima da série")
    estimativa = cesta.recalcular_e_gravar(pesquisa, item, repo,
                                           perfil=perfil.LEI_14133)
    assert estimativa.valor_unitario == Decimal("22.00")
    assert estimativa.metodo == "media"

    [item] = repo.listar_itens(salvo["pesquisa_id"])
    assert Decimal(str(item["preco_estimado"])) == Decimal("22.00")
    assert Decimal(str(item["preco_total"])) == Decimal("2200.00")
    repo.confirmar_item(str(item["id"]))
    [item] = repo.listar_itens(salvo["pesquisa_id"])
    assert item["estado"] == EstadoItem.COMPLETO.value

    # relatórios pelo mesmo gerador da pesquisa automática
    referencias = {str(item["id"]): repo.listar_referencias(str(item["id"]))}
    completo = relatorio.completo(pesquisa, [item], referencias)
    assert "22,00" in completo and "preço 9× acima da série" in completo
    assert "22,00" in relatorio.resumido(pesquisa, [item], referencias)

    from openpyxl import load_workbook

    planilha = load_workbook(BytesIO(
        relatorio.xlsx_analitico(pesquisa, [item], referencias)))
    linhas_refs = sum(1 for _ in planilha.worksheets[-1].iter_rows()) - 1
    assert linhas_refs >= 4, "a desconsiderada também vai para a memória"

    pdf = export.gerar_pdf("Relatório de Pesquisa de Preços", completo)
    assert pdf[:5] == b"%PDF-"

    # retorno ao processo pelo serviço de aplicação existente
    dados = {"itens": [{"descricao": ITEM["descricao"], "quantidade": 100,
                        "unidade": "UNIDADE", "valor_unitario": 1.0}]}
    novos, mudancas, recusas = aplicacao.aplicar(dados, pesquisa, [item])
    assert not recusas and len(mudancas) == 1
    assert Decimal(str(novos["itens"][0]["valor_unitario"])) == Decimal("22.00")
    assert novos[aplicacao.CHAVE_PROVENIENCIA]["metodologia"] == "media"
    assert cliente.tabelas["pesquisas_preco"][0]["id"] == salvo["pesquisa_id"]


def test_mesma_chave_nao_cria_segunda_pesquisa(monkeypatch):
    cliente, coleta, salvo = _salvar(monkeypatch)
    de_novo = cesta.salvar_consulta_como_pesquisa(
        repositorio=repo, nome="duplo clique", item=ITEM, coleta=coleta,
        selecionados=set(), perfil=perfil.LEI_14133, auth_user_id="u1",
        idempotency_key="banco:t1")
    assert de_novo["reaproveitada"] and de_novo["pesquisa_id"] == salvo["pesquisa_id"]
    assert len(cliente.tabelas["pesquisas_preco"]) == 1


def test_menos_de_tres_selecionadas_fica_parcial_e_incompleto(monkeypatch):
    cliente = ClientePrecosFalso()
    monkeypatch.setattr(db, "cliente_do_usuario", lambda: cliente)
    coleta = busca.coletar(dict(ITEM, id="c", numero=1), [_fonte()])
    duas = {r.chave_dedupe for r, _ in coleta.referencias[:2]}
    salvo = cesta.salvar_consulta_como_pesquisa(
        repositorio=repo, nome="duas", item=ITEM, coleta=coleta,
        selecionados=duas, perfil=perfil.LEI_14133, auth_user_id="u1")
    [item] = repo.listar_itens(salvo["pesquisa_id"])
    assert item["estado"] == EstadoItem.INCOMPLETO.value
    assert repo.obter_pesquisa(salvo["pesquisa_id"])["estado"] == "partial"
    # o valor provisório fica na memória, mas não chega ao processo:
    # só item CONCLUÍDO é aplicado
    dados = {"itens": [{"descricao": ITEM["descricao"], "quantidade": 100,
                        "unidade": "UNIDADE", "valor_unitario": 1.0}]}
    pesquisa = repo.obter_pesquisa(salvo["pesquisa_id"])
    novos, mudancas, _ = aplicacao.aplicar(dados, pesquisa, [item])
    assert all(m.unitario_novo is None for m in mudancas)
    assert novos["itens"][0]["valor_unitario"] == 1.0


def test_valor_estimado_marcado_na_tela_nao_entra_na_cesta(monkeypatch):
    """A tela manda chaves; o servidor confere a natureza (SRC-02)."""
    cliente = ClientePrecosFalso()
    monkeypatch.setattr(db, "cliente_do_usuario", lambda: cliente)
    estimado = _homologada("est", "5.00", NaturezaValor.ESTIMADO_ORIGEM)
    coleta = busca.coletar(dict(ITEM, id="c", numero=1), [_fonte(estimado)])
    assert estimado.chave_dedupe not in coleta.pre_selecao(Decimal("0"))
    todas = {r.chave_dedupe for r, _ in coleta.referencias}
    salvo = cesta.salvar_consulta_como_pesquisa(
        repositorio=repo, nome="com estimado", item=ITEM, coleta=coleta,
        selecionados=todas, perfil=perfil.LEI_14133, auth_user_id="u1")
    refs = cliente.tabelas["pesquisa_preco_referencias"]
    [linha] = [r for r in refs if r["id_externo"] == "est"]
    assert linha["status"] == "candidate"
    [item] = repo.listar_itens(salvo["pesquisa_id"])
    # 20/22/24/200 → 66,50; o 5,00 estimado ficou fora
    assert Decimal(str(item["preco_estimado"])) == Decimal("66.50")
