"""
G1 — confiabilidade do núcleo de Cotações.

Provas de aceite do G1 (`docs/cotacoes/gates.md`): exceção formal à
regra dos três, método por item, importação diagnosticada e idempotente,
a MESMA contagem de considerados em toda superfície, pendências nominais
no relatório e identificador que cobre o conteúdo decisório. Banco falso
das provas de preços; nenhuma rede, nenhuma IA.
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from src.precos import (cesta, estatistica, importacao, orientacao, perfil,
                        relatorio, unidades)
from src.precos import repositorio as repo
from src.precos.estados import EstadoItem
from tests.cotacoes_apoio import (cotacao_com_item, ligar_banco_falso, linha,
                                  referencia)


def _item_incompleto_com(valores, monkeypatch, perfil_id="lei_14133"):
    ligar_banco_falso(monkeypatch)
    pesquisa, item = cotacao_com_item(perfil_id=perfil_id)
    repo.registrar_referencias(str(item["id"]), [
        referencia(v, f"r{i}", normalizada=True) for i, v in enumerate(valores)])
    for r in repo.listar_referencias(str(item["id"])):
        repo.reclassificar_referencia(str(r["id"]), "selected", "escolhida")
    repo.mover_item(str(item["id"]), EstadoItem.BUSCANDO, EstadoItem.PENDENTE)
    repo.mover_item(str(item["id"]), EstadoItem.INCOMPLETO, EstadoItem.BUSCANDO)
    [item] = repo.listar_itens(str(pesquisa["id"]))
    cesta.recalcular_e_gravar(pesquisa, item, repo, perfil=perfil.obter(perfil_id))
    [item] = repo.listar_itens(str(pesquisa["id"]))
    return pesquisa, item


def _excecao(pesquisa, item, **extra):
    argumentos = dict(perfil=perfil.LEI_14133,
                      texto="Só duas contratações comparáveis no período.",
                      fundamento="mercado_restrito",
                      aprovador="Maria Souza, Secretária de Administração",
                      ator="u1")
    argumentos.update(extra)
    return cesta.registrar_excecao(pesquisa, item, repo, **argumentos)


# ---------------------------------------------------------------------------
# Exceção à regra dos três
# ---------------------------------------------------------------------------
def test_excecao_exige_perfil_preco_texto_aprovador_e_declaracao(monkeypatch):
    pesquisa, item = _item_incompleto_com(["29.00", "22.70"], monkeypatch)
    sem_excecao = perfil.PerfilNormativo(
        id="restrito", nome="Restrito", base_legal="x",
        admite_menos_com_justificativa=False)
    casos = [
        (dict(perfil=sem_excecao), "não admite"),
        (dict(texto="curto"), "15 caracteres"),
        (dict(aprovador=""), "autoridade competente"),
        (dict(aprovacao_declarada=False), "consta dos autos"),
        (dict(fundamento="inventado"), "fundamento"),
    ]
    for extra, esperado in casos:
        with pytest.raises(cesta.ExcecaoInvalida, match=esperado):
            _excecao(pesquisa, item, **extra)
    [item] = repo.listar_itens(str(pesquisa["id"]))
    assert item["estado"] == EstadoItem.INCOMPLETO.value


def test_excecao_nao_forma_preco_de_item_sem_preco(monkeypatch):
    pesquisa, item = _item_incompleto_com([], monkeypatch)
    with pytest.raises(cesta.ExcecaoInvalida, match="nenhum preço"):
        _excecao(pesquisa, item)


def test_excecao_leva_a_revisao_e_o_revisor_conclui(monkeypatch):
    pesquisa, item = _item_incompleto_com(["29.00", "22.70"], monkeypatch)
    est = _excecao(pesquisa, item)
    assert est.concluida and est.excecao_aplicada
    assert est.valor_unitario == Decimal("25.85")
    [item] = repo.listar_itens(str(pesquisa["id"]))
    assert item["estado"] == EstadoItem.EM_REVISAO.value
    assert "EXCEÇÃO à regra dos 3" in item["justificativa"]
    assert item["estatisticas"]["decisoes"]["excecao"]["consideradas"] == 2

    repo.confirmar_item(str(item["id"]), justificativa="conferido")
    [item] = repo.listar_itens(str(pesquisa["id"]))
    assert item["estado"] == EstadoItem.COMPLETO.value
    assert "EXCEÇÃO" in item["justificativa"]
    assert "conclusão do revisor: conferido" in item["justificativa"]

    eventos = [e["tipo"] for e in
               ligar_banco_falso.__globals__["db"].cliente_do_usuario()
               .tabelas["pesquisa_preco_eventos"]]
    assert "excecao_registrada" in eventos


def test_cesta_que_muda_invalida_a_excecao(monkeypatch):
    pesquisa, item = _item_incompleto_com(["29.00", "22.70"], monkeypatch)
    _excecao(pesquisa, item)
    [item] = repo.listar_itens(str(pesquisa["id"]))
    [primeira, _] = repo.listar_referencias(str(item["id"]))
    repo.reclassificar_referencia(str(primeira["id"]), "rejected", "erro")
    est = cesta.recalcular_e_gravar(pesquisa, item, repo, perfil=perfil.LEI_14133)
    assert not est.concluida
    [item] = repo.listar_itens(str(pesquisa["id"]))
    assert item["estado"] == EstadoItem.INCOMPLETO.value
    decisoes = item["estatisticas"]["decisoes"]
    assert "excecao" not in decisoes
    assert "mudou de 2 para 1" in decisoes["excecoes_anteriores"][0][
        "motivo_invalidacao"]


def test_excecao_em_lote_mostra_previa_antes_de_gravar(monkeypatch):
    pesquisa, item = _item_incompleto_com(["29.00", "22.70"], monkeypatch)
    [outro] = repo.salvar_itens(str(pesquisa["id"]), [{
        "numero": 2, "descricao": "CANETA", "unidade": "UNIDADE",
        "quantidade": "5"}])
    itens = repo.listar_itens(str(pesquisa["id"]))
    referencias = {str(i["id"]): repo.listar_referencias(str(i["id"]))
                   for i in itens}
    previa = cesta.prever_excecao_em_lote(itens, referencias,
                                          perfil=perfil.LEI_14133,
                                          escopo="todos")
    por_numero = {p["numero"]: p for p in previa}
    assert por_numero[1]["aplica"] is True
    assert por_numero[2]["aplica"] is False
    assert "sem preço" in por_numero[2]["situacao"]
    # nada gravado pela prévia
    assert all("decisoes" not in (i.get("estatisticas") or {})
               for i in repo.listar_itens(str(pesquisa["id"])))
    resultado = cesta.aplicar_excecao_em_lote(
        pesquisa, itens, repo, previa=previa, perfil=perfil.LEI_14133,
        fundamento="mercado_restrito",
        texto_modelo=cesta.modelo_de_justificativa("mercado_restrito", 2),
        aprovador="Autoridade X")
    assert [r for _, r in resultado] == ["exceção registrada"]


# ---------------------------------------------------------------------------
# Método do item
# ---------------------------------------------------------------------------
def test_metodo_do_revisor_persiste_entre_recalculos(monkeypatch):
    pesquisa, item = _item_incompleto_com(["10.00", "11.00", "30.00"],
                                          monkeypatch)
    with pytest.raises(ValueError, match="justifique"):
        cesta.escolher_metodo(pesquisa, item, repo, perfil=perfil.LEI_14133,
                              metodo="mediana", justificativa="x")
    est = cesta.escolher_metodo(
        pesquisa, item, repo, perfil=perfil.LEI_14133, metodo="mediana",
        justificativa="série com valor muito acima do mercado local")
    assert est.metodo == "mediana" and est.valor_unitario == Decimal("11.00")
    [item] = repo.listar_itens(str(pesquisa["id"]))
    de_novo = cesta.recalcular_e_gravar(pesquisa, item, repo,
                                        perfil=perfil.LEI_14133)
    assert de_novo.metodo == "mediana"
    est = cesta.escolher_metodo(pesquisa, item, repo, perfil=perfil.LEI_14133,
                                metodo="automatico", justificativa="")
    assert est.metodo == "media" and est.valor_unitario == Decimal("17.00")


# ---------------------------------------------------------------------------
# Unidades
# ---------------------------------------------------------------------------
def test_dicionario_novo_nao_converte_entre_unidades_distintas():
    assert unidades.canonizar("PAC") == "PACOTE"
    assert unidades.canonizar("pc.") == "UNIDADE"
    assert unidades.canonizar("M²") == "METRO QUADRADO"
    assert unidades.canonizar("Caixotes") is None
    assert unidades.canonizar("UNDADE") is None
    fator, _ = unidades.fator_de_conversao("SACO", "QUILOGRAMA", None)
    assert fator is None


# ---------------------------------------------------------------------------
# Importação
# ---------------------------------------------------------------------------
def _xlsx(linhas, cabecalho=("Código", "Descrição", "Unidade", "Quantidade",
                             "UF", "Lote")) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(list(cabecalho))
    for linha in linhas:
        ws.append(list(linha))
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _linhas_210():
    return [(f"C{n:03d}", f"ITEM SINTETICO {n}", "UNIDADE", n, "PA",
             f"L{n % 3}") for n in range(1, 211)]


def test_importacao_de_210_itens_preserva_codigos_e_e_idempotente(monkeypatch):
    ligar_banco_falso(monkeypatch)
    pesquisa, _ = cotacao_com_item()
    pid = str(pesquisa["id"])
    diag = importacao.diagnosticar(_xlsx(_linhas_210()))
    assert diag.pode_importar and len(diag.validos) == 210
    plano = importacao.planejar(diag, repo.listar_itens(pid))
    assert importacao.resumo_do_plano(plano)["novo"] == 210
    importacao.aplicar(pid, plano, repo)
    itens = repo.listar_itens(pid)
    assert len(itens) == 211
    assert {i["codigo"] for i in itens if i["codigo"]} == \
        {f"C{n:03d}" for n in range(1, 211)}
    assert all(i.get("uf") == "PA" for i in itens if i["codigo"])

    # reimportar a MESMA planilha, embaralhada: nada muda
    embaralhada = list(reversed(_linhas_210()))
    plano = importacao.planejar(importacao.diagnosticar(_xlsx(embaralhada)),
                                repo.listar_itens(pid))
    assert importacao.resumo_do_plano(plano) == {
        "novo": 0, "atualizar": 0, "igual": 210, "refaz_busca": 0}
    importacao.aplicar(pid, plano, repo)
    assert len(repo.listar_itens(pid)) == 211


def test_reimportar_nao_perde_selecao_e_unidade_mudada_refaz_busca(monkeypatch):
    pesquisa, item = _item_incompleto_com(["29.00", "22.70"], monkeypatch)
    pid = str(pesquisa["id"])
    repo.atualizar_pesquisa(pid, nome="com código")
    repo.salvar_itens(pid, [{"numero": 1, "codigo": "P1",
                             "descricao": item["descricao"],
                             "unidade": "RESMA", "quantidade": "10"}])
    antes = repo.listar_referencias(str(item["id"]))
    # mesma linha → igual; estado e seleção intocados
    plano = importacao.planejar(importacao.diagnosticar(_xlsx(
        [("P1", item["descricao"], "RESMA", 10, "", "")])), repo.listar_itens(pid))
    assert plano[0].tipo == "igual"
    importacao.aplicar(pid, plano, repo)
    [atual] = repo.listar_itens(pid)
    assert atual["estado"] == EstadoItem.INCOMPLETO.value
    assert repo.listar_referencias(str(item["id"])) == antes
    # unidade trocada num item já pesquisado → volta à busca
    plano = importacao.planejar(importacao.diagnosticar(_xlsx(
        [("P1", item["descricao"], "PACOTE", 10, "", "")])), repo.listar_itens(pid))
    assert plano[0].tipo == "atualizar" and plano[0].refaz_busca
    importacao.aplicar(pid, plano, repo)
    [atual] = repo.listar_itens(pid)
    assert atual["estado"] == EstadoItem.BUSCANDO.value
    assert atual["unidade"] == "PACOTE"


def test_importacao_com_erro_nao_grava_nada(monkeypatch):
    diag = importacao.diagnosticar(_xlsx([
        ("A", "Caneta", "UNIDADE", 1, "PA", ""),
        ("A", "Lápis", "UND", "1,5", "PA", ""),
        ("", "Régua", "UNIDADE", 0, "", ""),
        ("", "", "", "", "", ""),
    ]))
    assert not diag.pode_importar
    assert diag.linhas_com_erro == [3, 4]
    motivos = " ".join(e.motivo for e in diag.erros)
    assert "repetido" in motivos and "maior que zero" in motivos
    assert diag.em_branco == 1
    assert importacao.quantidade_de("1.234,5") == Decimal("1234.5")
    assert importacao.quantidade_de("12.5") == Decimal("12.5")
    assert importacao.quantidade_de("abc") is None


# ---------------------------------------------------------------------------
# A mesma contagem em todas as superfícies
# ---------------------------------------------------------------------------
def test_considerados_sao_os_mesmos_em_toda_superficie():
    linhas = [linha("10.00"), linha("11.00"), linha("30.00", "warning"),
              linha("12.00", "selected", natureza="estimado_origem"),
              linha("13.00", "rejected", motivos=["fora da curva"])]
    linhas[3]["valor_unitario_normalizado"] = "12.00"
    item = {"id": "item-1", "numero": 1, "descricao": "PAPEL",
            "unidade": "RESMA", "estado": "review", "quantidade": "1",
            "estatisticas": {}}
    assert cesta.validas_na_cesta(linhas) == 3
    est = cesta.recalcular(linhas, perfil=perfil.LEI_14133,
                           regras=estatistica.REGRAS_MEDIA_PADRAO)
    assert len(est.cesta.selecionadas) == 3
    assert est.valor_unitario == Decimal("17.00")
    pesquisa = {"id": "p", "nome": "t"}
    completo = relatorio.completo(pesquisa, [item], {"item-1": linhas})
    assert "consideradas no cálculo (3)" in completo
    assert "Referências não consideradas (2)" in completo
    assert "marcada, mas fora do cálculo" in completo
    resumido = relatorio.resumido(pesquisa, [item], {"item-1": linhas})
    assert "| 3/3 |" in resumido
    wb = load_workbook(BytesIO(relatorio.xlsx_analitico(
        pesquisa, [item], {"item-1": linhas})))
    cabecalho = [c.value for c in wb["Itens"][1]]
    valor = wb["Itens"][2][cabecalho.index("Consideradas no cálculo")].value
    assert valor == 3
    avisos = orientacao.do_item(item, linhas, perfil.LEI_14133)
    assert not any("abaixo do mínimo" in a.texto for a in avisos)


# ---------------------------------------------------------------------------
# Relatório: pendências nominais, exceção impressa, citação, hash
# ---------------------------------------------------------------------------
def test_pendencias_nominais_dizem_qual_item_e_qual_deficiencia():
    itens = [
        {"id": "a", "numero": 1, "descricao": "PAPEL", "estado": "complete"},
        {"id": "b", "numero": 2, "descricao": "CANETA", "estado": "incomplete",
         "estatisticas": {}},
        {"id": "c", "numero": 3, "descricao": "LÁPIS", "estado": "error"},
        {"id": "d", "numero": 4, "descricao": "RÉGUA", "estado": "review",
         "estatisticas": {"pendencias": ["dispersão elevada"]}},
    ]
    refs = {"b": [linha("5.00")], "d": [linha("1"), linha("2"), linha("3")]}
    nominais = relatorio.pendencias_nominais(itens, refs, perfil.LEI_14133)
    por_numero = {p["numero"]: p for p in nominais}
    assert set(por_numero) == {2, 3, 4}
    assert "1/3 preços considerados" in por_numero[2]["deficiencias"][0]
    assert "não responderam" in por_numero[3]["deficiencias"][0]
    assert any("dispersão" in d for d in por_numero[4]["deficiencias"])
    texto = relatorio.completo({"id": "p"}, itens, refs)
    assert "**Pendências por item**" in texto and "CANETA" in texto


def test_relatorio_imprime_excecao_e_cita_a_norma_do_perfil(monkeypatch):
    pesquisa, item = _item_incompleto_com(["29.00", "22.70", "9.66"],
                                          monkeypatch, perfil_id="in_65_2021")
    [terceira] = [r for r in repo.listar_referencias(str(item["id"]))
                  if Decimal(str(r["valor_unitario_normalizado"])) == Decimal("9.66")]
    repo.reclassificar_referencia(str(terceira["id"]), "rejected",
                                  "inexequível frente aos demais")
    [item] = repo.listar_itens(str(pesquisa["id"]))
    cesta.recalcular_e_gravar(pesquisa, item, repo, perfil=perfil.IN_65_2021)
    [item] = repo.listar_itens(str(pesquisa["id"]))
    _excecao(pesquisa, item, perfil=perfil.IN_65_2021)
    [item] = repo.listar_itens(str(pesquisa["id"]))
    refs = {str(item["id"]): repo.listar_referencias(str(item["id"]))}
    texto = relatorio.completo(pesquisa, [item], refs)
    assert "IN SEGES/ME nº 65/2021, art. 3º, VI" in texto
    assert "IN SEGES/ME nº 65/2021, art. 6º, § 5º" in texto
    assert "Maria Souza" in texto and "R$ 25,85" in texto


def test_identificador_v1_continua_conferindo_relatorios_antigos():
    """Hash calculado com o código do commit 0de2939 (formato 1)."""
    pesquisa = {"id": "p-1", "versao": 2, "perfil_normativo": "in_65_2021",
                "data_base": "2026-10-01"}
    itens = [{"id": "i1", "numero": 1, "descricao": "PAPEL A4",
              "estado": "complete", "metodo": "media",
              "preco_estimado": "20.45"}]
    refs = {"i1": [{"raw_hash": "abc", "status": "selected"},
                   {"raw_hash": "def", "status": "rejected"}]}
    assert relatorio.identificador_da_versao_v1(pesquisa, itens, refs) == \
        "8bbe10b6969e155be4b4724ba823a5bdd9e5b422785c9c896f206c41257f65ea"
    assert relatorio.identificador_da_versao(pesquisa, itens, refs) != \
        relatorio.identificador_da_versao_v1(pesquisa, itens, refs)


def test_identificador_muda_com_motivo_e_com_decisao_do_revisor():
    item = {"id": "i", "numero": 1, "estado": "complete",
            "estatisticas": {"decisoes": {}}}
    refs = {"i": [linha("10", "rejected", motivos=["a"])]}
    base = relatorio.identificador_da_versao({"id": "p"}, [item], refs)
    refs2 = {"i": [linha("10", "rejected", motivos=["b"])]}
    assert relatorio.identificador_da_versao({"id": "p"}, [item], refs2) != base
    item2 = {**item, "estatisticas": {"decisoes": {"metodo": {"metodo": "mediana"}}}}
    assert relatorio.identificador_da_versao({"id": "p"}, [item2], refs) != base


# ---------------------------------------------------------------------------
# Critério percentual configurável na estimativa
# ---------------------------------------------------------------------------
def test_criterio_percentual_da_cotacao_gera_pendencia_e_memoria():
    linhas = [linha("29.00"), linha("22.70"), linha("9.66")]
    est = cesta.recalcular(
        linhas, perfil=perfil.LEI_14133,
        regras=estatistica.REGRAS_MEDIA_PADRAO,
        criterio_percentual=estatistica.criterio_percentual_da(
            {"criterio_percentual": {"elevado": "30", "inexequivel": "70"}}))
    assert est.valor_unitario == Decimal("20.45")
    assert any("42,56% da média dos outros" in m for m in est.memoria)
    assert any("critério percentual" in p for p in est.pendencias)
    gravado = est.para_relatorio()
    assert gravado["criterio_percentual"]["elevado"] == "30"
    assert estatistica.criterio_percentual_da({}) is None
    assert estatistica.criterio_percentual_da(
        {"criterio_percentual": {"elevado": "30", "inexequivel": "170"}}) is None
