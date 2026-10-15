"""Generate synthetic fidelity inputs with the existing product exporter.

No tenant documents, network, database, AI calls or second export engine.
Decimal is an independent fixture oracle, not a new production calculator.
"""
from __future__ import annotations

import argparse
import base64
import copy
from datetime import date
from decimal import Decimal
import hashlib
import io
import json
from pathlib import Path
import sys
import zipfile
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src import planilha
from src.export import gerar_docx, gerar_pdf, motor_pdf

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def itens_sinteticos() -> list[dict]:
    return [{"codigo": f"GC-SYN-{i:04d}",
             "descricao": f"Material sintético {i:03d} para ensaio de tabela e paginação",
             "unidade": "UN", "quantidade": str(1 + i % 7),
             "valor_unitario": f"{100 + i // 10},{i % 100:02d}",
             "fonte": f"https://example.invalid/precos/{i}"} for i in range(1, 211)]


def tabela_sintetica() -> tuple[str, dict]:
    source = itens_sinteticos()
    expected = [Decimal(row["quantidade"]) * Decimal(row["valor_unitario"].replace(",", "."))
                for row in source]
    computed, total = planilha.calcular(source)
    if len(computed) != 210 or any(Decimal(str(row["valor_total"])) != wanted
                                  for row, wanted in zip(computed, expected)):
        raise ValueError("Existing calculator differs from the independent fixture oracle")
    if Decimal(str(total)) != sum(expected):
        raise ValueError("Existing global total differs from the fixture oracle")
    return planilha.para_markdown(computed, total), {
        "items": 210, "codes": [row["codigo"] for row in source],
        "line_totals_decimal": [str(value.quantize(Decimal("0.01"))) for value in expected],
        "global_total_decimal": str(sum(expected).quantize(Decimal("0.01")))}


def inspect_docx(data: bytes) -> dict:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
        page = root.find(f".//{W}pgSz")
        text = "\n".join(node.text or "" for node in root.iter(W + "t"))
        return {"page_twips": {key.rsplit("}", 1)[-1]: value for key, value in page.attrib.items()},
                "table_rows": [len(table.findall(W + "tr")) for table in root.iter(W + "tbl")],
                "hyperlinks": len(list(root.iter(W + "hyperlink"))),
                "headers": len([name for name in archive.namelist() if name.startswith("word/header") and name.endswith(".xml")]),
                "footers": len([name for name in archive.namelist() if name.startswith("word/footer") and name.endswith(".xml")]),
                "text": text}


def text_blocks(data: bytes) -> dict:
    """Read semantic text boundaries, ignoring run splits and layout whitespace."""
    result = {"body": [], "headers": [], "footers": [], "links": []}
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for name in archive.namelist():
            group = ("body" if name == "word/document.xml" else
                     "headers" if name.startswith("word/header") and name.endswith(".xml") else
                     "footers" if name.startswith("word/footer") and name.endswith(".xml") else None)
            if group:
                root = ET.fromstring(archive.read(name))
                blocks = [" ".join("".join(t.text or "" for t in p.iter(W + "t")).split()) for p in root.iter(W + "p")]
                result[group].extend(block for block in blocks if block)
            if name.startswith("word/_rels/") and name.endswith(".rels"):
                for relation in ET.fromstring(archive.read(name)):
                    if relation.get("Type", "").endswith("/hyperlink"):
                        result["links"].append(relation.get("Target"))
    for group in ("headers", "footers", "links"):
        result[group].sort()
    return result


def verify_inputs(inputs: Path) -> dict:
    manifest = json.loads((inputs / "inputs.json").read_text(encoding="utf-8"))
    if manifest.get("synthetic_only") is not True or len(manifest["files"]) != 3:
        raise ValueError("Only the reviewed synthetic corpus can enter this QA path")
    for item in manifest["files"]:
        name = item["path"]
        if Path(name).name != name or "\\" in name or ":" in name or not name.endswith(".docx"):
            raise ValueError("Invalid synthetic fixture path")
        data = (inputs / name).read_bytes()
        if len(data) > 8 * 1024 * 1024 or hashlib.sha256(data).hexdigest() != item["sha256"]:
            raise ValueError("Synthetic fixture digest differs")
    return manifest


def compare_roundtrip(inputs: Path, output: Path) -> dict:
    manifest = verify_inputs(inputs)
    report = {"status": "SEMANTIC_ROUNDTRIP_CHECKED_VISUAL_REVIEW_PENDING", "cases": []}
    for item in manifest["files"]:
        name = Path(item["path"]).stem
        before = text_blocks((inputs / item["path"]).read_bytes())
        after = text_blocks((output / name / "roundtrip.docx").read_bytes())
        if before != after:
            raise ValueError(f"WordCraft roundtrip changed semantic text or links: {name}")
        if name.startswith("02"):
            model_text = (output / name / "model.txt").read_text(encoding="utf-8")
            if any(code not in model_text for code in manifest["oracle"]["codes"]):
                raise ValueError("Native model dropped a synthetic item")
        report["cases"].append({"name": name, "semantic_text_and_links": "PASS", "visual_review": "PENDING"})
    (output / "comparison.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    return report


def generate(output: Path, *, pdf=False) -> dict:
    if pdf and motor_pdf() != "libreoffice":
        raise ValueError("Visual reference requires the real institutional LibreOffice engine")
    table, oracle = tabela_sintetica()
    intro = ("Corpus inteiramente sintético para comparar importação, edição e exportação. "
             "Não representa processo administrativo, órgão ou assinatura real.\n\n")
    signatures = ("\n\n## Assinaturas de teste\n\n**Pessoa Sintética A**\n\n"
                  "Responsável fictício pelo ensaio\n\n**Pessoa Sintética B**\n\nRevisor fictício\n")
    cases = {
        "01-texto-listas-links": ("Ensaio de texto listas e links", intro +
            "# Objeto sintético\n\nParágrafo com **negrito**, *itálico* e caracteres ç ã é §.\n\n"
            "## Requisitos\n\n- Preservar primeiro item\n- Preservar segundo item\n\n"
            "1. Etapa inicial\n2. Etapa final\n\n[Referência sintética](https://example.invalid/referencia)\n", None, ""),
        "02-tabela-210-itens": ("Ensaio da planilha com 210 itens", intro + table, None, ""),
        "03-timbrado-assinaturas": ("Ensaio de identidade e assinaturas", intro +
            "# Identidade institucional sintética\n\nO cabeçalho e o rodapé devem permanecer na página A4.\n\n"
            "## Quadro de teste\n\n| Campo | Valor |\n|---|---|\n| Revisão | Sintética 1 |\n| Escopo | Ensaio local |\n",
            {"cabecalho": "ÓRGÃO SINTÉTICO PARA TESTES DE FIDELIDADE",
             "rodape": "Ensaio GovConnect sem validade administrativa"}, signatures),
    }
    output = output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    manifest = {"schema": "gc-wordcraft-corpus/v1", "synthetic_only": True,
                "status": "INPUTS_GENERATED_NOT_ROUNDTRIP_APPROVED", "oracle": oracle, "cases": []}
    for name, (title, markdown, branding, signature) in cases.items():
        data = gerar_docx(title, markdown, branding, signature)
        inspection = inspect_docx(data)
        if inspection["page_twips"] != {"w": "11906", "h": "16838"}:
            raise ValueError("Existing export is not the expected A4 fixture")
        if name.startswith("02") and any(code not in inspection["text"] for code in oracle["codes"]):
            raise ValueError("The exported fixture dropped an item")
        (output / (name + ".docx")).write_bytes(data)
        (output / (name + ".md")).write_text(markdown + signature, encoding="utf-8")
        inspection.pop("text")
        entry = {"name": name, "sha256_docx": hashlib.sha256(data).hexdigest(), **inspection}
        if pdf:
            from pypdf import PdfReader
            rendered = gerar_pdf(title, markdown, branding, signature)
            if motor_pdf() != "libreoffice":
                raise ValueError("PDF fallback cannot approve the institutional reference")
            reader = PdfReader(io.BytesIO(rendered))
            extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
            if name.startswith("02") and any(code not in extracted for code in oracle["codes"]):
                raise ValueError("PDF reference dropped or split a synthetic item code")
            if name.startswith("03") and any(label not in extracted for label in
                    (branding["cabecalho"], branding["rodape"], "Pessoa Sintética A", "Pessoa Sintética B")):
                raise ValueError("PDF reference lost institutional text or signatures")
            (output / (name + ".pdf")).write_bytes(rendered)
            entry.update(sha256_pdf=hashlib.sha256(rendered).hexdigest(), pages=len(reader.pages),
                         pdf_engine=motor_pdf(), visual_review="PENDING")
        manifest["cases"].append(entry)
    (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


# Full G3 fixtures are deliberately separate from the three immutable G1 inputs.
FULL_SCHEMA = "gc-wordcraft-full-corpus/v1"
SYNTHETIC_NOTICE = (
    "Corpus inteiramente sintético e sem validade administrativa. "
    "Nomes, identidades, datas, portarias, preços e recomendações são dados de QA. "
    "Não houve consulta a banco, rede ou modelo de linguagem.\n\n")


def _moeda_oraculo(value: Decimal) -> str:
    text = f"{value.quantize(Decimal('0.01')):,.2f}"
    return "R$ " + text.replace(",", "_").replace(".", ",").replace("_", ".")


def tabela_com_oraculo(source: list[dict]) -> tuple[str, dict]:
    """Independent Decimal oracle first; production calculator/exporter second."""
    expected = [Decimal(row["quantidade"]) * Decimal(row["valor_unitario"].replace(",", "."))
                for row in source]
    total_expected = sum(expected, Decimal("0.00"))
    computed, total = planilha.calcular(copy.deepcopy(source))
    if (len(computed) != len(source)
            or any(Decimal(str(row["valor_total"])) != wanted for row, wanted in zip(computed, expected))
            or Decimal(str(total)) != total_expected):
        raise ValueError("Existing calculator differs from the independent full-corpus oracle")
    oracle = {
        "items": len(source), "codes": [row["codigo"] for row in source],
        "source_items": copy.deepcopy(source),
        "line_totals_decimal": [str(value.quantize(Decimal("0.01"))) for value in expected],
        "global_total_decimal": str(total_expected.quantize(Decimal("0.01"))),
        "expected_money": [_moeda_oraculo(value) for value in expected],
        "global_total_display": _moeda_oraculo(total_expected),
        "external_links": [row["fonte"] for row in source],
    }
    return planilha.para_markdown(computed, total), oracle


def _logo_sintetico() -> bytes:
    """Original geometric QA banner, encoded with stdlib; no borrowed assets."""
    import struct
    import zlib
    width, height = 600, 48
    rows = []
    for y in range(height):
        row = bytearray()
        for x in range(width):
            square = 12 <= x < 36 and 12 <= y < 36
            stripe = 56 <= x < 570 and 20 <= y < 28
            row.extend((255, 255, 255) if square or stripe else (24, 62, 86))
        rows.append(b"\0" + row)

    def chunk(name, value):
        return struct.pack(">I", len(value)) + name + value + struct.pack(
            ">I", zlib.crc32(name + value) & 0xffffffff)

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(b"".join(rows))) + chunk(b"IEND", b""))


def instituicao_sintetica(tag="A") -> dict:
    """Reuse only pure context/eligibility/snapshot rules; no fake DB approval."""
    from uuid import NAMESPACE_URL, uuid5
    from src import assinaturas, contexto, portarias
    identifier = lambda kind: str(uuid5(NAMESPACE_URL, "gc-full-corpus:" + tag + ":" + kind))
    secretary = {"id": identifier("secretaria"), "nome": "Secretaria Sintética " + tag,
                 "cabecalho": "SECRETARIA SINTÉTICA " + tag,
                 "rodape": "QA sem validade administrativa " + tag}
    municipal = {"id": identifier("municipio"), "nome": "Município Sintético " + tag,
                 "padrao": True, "cabecalho": "MUNICÍPIO SINTÉTICO " + tag,
                 "rodape": "Herança institucional sintética " + tag}
    identity, origin = contexto.resolver_identidade([municipal, secretary], secretary["id"])
    branding_snapshot = copy.deepcopy(identity)
    ordinance = {"id": identifier("portaria"), "secretaria_id": secretary["id"],
                 "tipo": portarias.EQUIPE_PLANEJAMENTO, "status": portarias.ATIVA,
                 "numero": "QA-001", "ano": 2026, "data_publicacao": "2026-01-05",
                 "data_inicio_vigencia": "2026-01-05", "data_fim_vigencia": "2026-12-31"}
    resolved = portarias.resolver([ordinance], secretaria_id=secretary["id"],
                                  data_referencia=date(2026, 10, 14))
    servers = [{"id": identifier("pessoa-" + str(i)), "nome": f"Pessoa Sintética {tag}{i}",
                "matricula": f"QA-{tag}-{i}", "cargo": "Cargo exclusivamente fictício", "ativo": True}
               for i in (1, 2)]
    members = [{"portaria_id": ordinance["id"], "servidor_id": server["id"],
                "funcao": "EQUIPE_PLANEJAMENTO", "ativo": True, "ordem": i}
               for i, server in enumerate(servers, 1)]
    eligible = assinaturas.elegiveis(funcao="EQUIPE_PLANEJAMENTO", servidores=servers,
        funcoes_do_servidor=[], portaria=resolved, membros_da_portaria=members,
        secretaria_id=secretary["id"], data_referencia=date(2026, 10, 14))
    snapshots = []
    for order, server in enumerate(eligible, 1):
        assinaturas.conferir_elegibilidade(servidor_id=server["id"],
            funcao="EQUIPE_PLANEJAMENTO", elegiveis_agora=eligible)
        snapshots.append(assinaturas.montar_snapshot(servidor=server,
            funcao="EQUIPE_PLANEJAMENTO", secretaria_nome=secretary["nome"],
            portaria=resolved, ordem=order))
    lines = assinaturas.linhas_de_exportacao(snapshots)
    signature = "\n\n# Assinaturas exclusivamente sintéticas\n\n" + "\n\n".join(
        "\n\n".join(block) for block in lines)
    return {"tenant_id": identifier("tenant"), "secretaria_id": secretary["id"],
            "origin": origin, "branding_snapshot": branding_snapshot,
            "municipal": municipal, "secretary": secretary, "signatarios_snapshot": snapshots,
            "portaria_snapshot": copy.deepcopy(resolved), "assinaturas_markdown": signature,
            "signature_names": [block[0] for block in lines],
            "approval_status": "FIXTURE_ONLY_NOT_DATABASE_APPROVAL"}


def _texto_perfil(doc_key, case_id, *, table="", long=False):
    from src import perfis
    clauses = perfis.clausulas_aplicaveis(doc_key, srp=True)
    pieces, markers, table_added = [SYNTHETIC_NOTICE], [], False
    for clause in clauses:
        pieces.append(f"## {clause['n']}. {clause['titulo']}\n")
        for number in range(1, 6 if long else 2):
            marker = f"{case_id}-S{clause['n']:02d}-P{number:02d}"
            markers.append(marker)
            pieces.append(
                f"{marker} O cenário fictício desta seção trata de {clause['titulo'].lower()}. "
                "A equipe de ensaio recebe quantidades declaradas e verifica os documentos de origem "
                "sem criar dados financeiros ou nomes de pessoas. Cada registro possui responsável "
                "sintético, prazo de teste e evidência identificável. A análise compara alternativas "
                "de organização, descreve critérios de recebimento e registra as dúvidas de execução. "
                "As conclusões pertencem apenas ao conjunto de dados fabricado para verificar "
                "paginação, acentuação, fluxo de parágrafos e continuidade dos títulos. Nenhuma "
                "passagem autoriza contratação, emissão, assinatura ou alteração de um processo real.\n")
        if table and clause.get("tabela") and not table_added:
            pieces.append(table + "\n")
            table_added = True
    return "\n".join(pieces), {
        "profile": doc_key, "clause_titles": [f"{c['n']}. {c['titulo']}" for c in clauses],
        "markers": markers}


def _enriquecer_qa(data: bytes, kind: str, *, clause_titles=()) -> bytes:
    """QA-only authoring of structures absent from the legacy Markdown surface."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document(io.BytesIO(data))
    if kind == "risk_merges":
        table = document.tables[0]
        table.cell(0, 0).merge(table.cell(0, 5))
        for start in (2, 6, 10):
            table.cell(start, 1).merge(table.cell(start + 3, 1))
        page = document.add_paragraph("Continuação da matriz de riscos sintética")
        page.paragraph_format.page_break_before = True
        table._tbl.addprevious(page._p)
    elif kind == "bookmarks":
        for number, title in enumerate(clause_titles, 1):
            matches = [p for p in document.paragraphs if p.text == title]
            if len(matches) != 1:
                raise ValueError("Canonical exporter lost a synthetic clause heading")
            paragraph = matches[0]
            start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
            start.set(qn("w:id"), str(number))
            start.set(qn("w:name"), f"gc_c05_{number:02d}")
            end.set(qn("w:id"), str(number))
            paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
            paragraph._p.append(end)
            reference = document.add_paragraph("Remissão sintética para a cláusula ")
            link = OxmlElement("w:hyperlink")
            link.set(qn("w:anchor"), f"gc_c05_{number:02d}")
            run, text = OxmlElement("w:r"), OxmlElement("w:t")
            text.text = str(number)
            run.append(text)
            link.append(run)
            reference._p.append(link)
    elif kind == "external_model":
        table = document.tables[0]
        table.cell(0, 0).merge(table.cell(0, 1))
        nested = table.cell(1, 1).add_table(rows=2, cols=2)
        nested.style = "Table Grid"
        for cell, value in zip((c for row in nested.rows for c in row.cells),
                               ("C10-NEST-A", "Campo", "C10-NEST-B", "Valor sintético")):
            cell.text = value
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        run, text = OxmlElement("w:r"), OxmlElement("w:t")
        text.text = "1"
        run.append(text)
        field.append(run)
        document.sections[0].footer.paragraphs[0]._p.append(field)
        document.add_page_break()
        document.add_paragraph("C10-PAGE-2 Continuação do modelo externo simulado")
    else:
        raise ValueError("Unknown QA fixture enrichment")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def inspect_full_docx(data: bytes) -> dict:
    inspection = inspect_docx(data)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
        inspection.update(
            bookmarks=[node.get(W + "name") for node in root.iter(W + "bookmarkStart")],
            internal_links=[node.get(W + "anchor") for node in root.iter(W + "hyperlink")
                            if node.get(W + "anchor")],
            horizontal_spans=[int(node.get(W + "val")) for node in root.iter(W + "gridSpan")],
            vertical_merge_restarts=sum(node.get(W + "val") == "restart" for node in root.iter(W + "vMerge")),
            vertical_merge_continuations=sum(node.get(W + "val") in {None, "continue"} for node in root.iter(W + "vMerge")),
            page_break_before=len(list(root.iter(W + "pageBreakBefore"))),
            explicit_page_breaks=sum(node.get(W + "type") == "page" for node in root.iter(W + "br")),
            media_sha256=sorted(hashlib.sha256(archive.read(name)).hexdigest()
                                for name in archive.namelist() if name.startswith("word/media/")))
    return inspection


def full_corpus_cases() -> list[dict]:
    """Literal synthetic scenarios and independent expectations, before export."""
    from src import blocos, config, contexto
    institution = instituicao_sintetica()
    branding = institution["branding_snapshot"]
    signature = institution["assinaturas_markdown"]

    def variant(name, title, markdown, *, oracle=None, qa=None, brand=None, sign=""):
        return {"name": name, "title": title, "markdown": markdown,
                "oracle": oracle or {}, "qa": qa, "branding": copy.deepcopy(brand),
                "signatures": sign}

    def case(case_id, name, variants, *, workflow=None, support=None):
        return {"id": case_id, "name": name, "variants": variants,
                "workflow_acceptance": workflow or [], "support": support or {}}

    dfd_rows = [{"codigo": f"GC-C01-{i:04d}", "descricao": f"Insumo sintético de planejamento {i}",
                 "unidade": "UN", "quantidade": str(i), "valor_unitario": price,
                 "fonte": f"https://example.invalid/c01/{i}"}
                for i, price in enumerate(("0,10", "0,20", "0,30", "1250,55", "999,99"), 1)]
    dfd_table, dfd_oracle = tabela_com_oraculo(dfd_rows)
    dfd_money_oracle = copy.deepcopy(dfd_oracle)
    dfd, dfd_profile = _texto_perfil("dfd", "C01", table=dfd_table)
    dfd_oracle.update(dfd_profile)
    dfd_oracle.update(signature_names=institution["signature_names"],
                      signatarios_snapshot=institution["signatarios_snapshot"],
                      portaria_snapshot=institution["portaria_snapshot"],
                      institutional_origin=institution["origin"])
    etp, etp_oracle = _texto_perfil("etp", "C02", long=True)
    tr_rows = [{"codigo": f"GC-C04-{i:04d}", "descricao": f"Material sintético de referência {i:02d}",
                "unidade": "UN", "quantidade": str(i % 5 + 1),
                "valor_unitario": str(Decimal("10.05") + Decimal(i) * Decimal("0.13")).replace(".", ","),
                "fonte": f"https://example.invalid/c04/{i}"} for i in range(1, 31)]
    tr_table, tr_oracle = tabela_com_oraculo(tr_rows)
    tr, tr_profile = _texto_perfil("tr", "C04", table=tr_table)
    tr_oracle.update(tr_profile)

    risk_codes = [f"C03-R{i:02d}" for i in range(1, 13)]
    risk_lines = ["| Matriz sintética de riscos | | | | | |", "|---|---|---|---|---|---|",
                  "| Código | Categoria | Evento | Probabilidade | Impacto | Tratamento |"]
    for i, code in enumerate(risk_codes):
        category = ("Planejamento", "Seleção", "Execução")[i // 4] if i % 4 == 0 else ""
        risk_lines.append(
            f"| {code} | {category} | Evento fictício {i + 1} com descrição integral para "
            "verificar quebra de linha e leitura da célula. | Média | Moderado | "
            "Conferência sintética da evidência, registro da dúvida e revisão pelo responsável de QA. |")
    risks = SYNTHETIC_NOTICE + "## Matriz de riscos de ensaio\n\n" + "\n".join(risk_lines)
    edital_titles = [f"{i}. {title}" for i, title in enumerate((
        "OBJETO SINTÉTICO", "CONDIÇÕES DE PARTICIPAÇÃO", "DADOS DO CENÁRIO", "ENVIO DAS PROPOSTAS",
        "DESCRIÇÃO DOS ITENS", "PREÇOS DE TESTE", "PRAZOS DE ENSAIO", "ORGANIZAÇÃO DA SESSÃO",
        "CRITÉRIOS DE COMPARAÇÃO", "AVALIAÇÃO DAS AMOSTRAS", "TRATAMENTO DE DÚVIDAS", "REGISTRO DAS EVIDÊNCIAS",
        "DOCUMENTOS DO FORNECEDOR FICTÍCIO", "VERIFICAÇÃO DOS REGISTROS", "COMUNICAÇÕES DE TESTE",
        "REVISÃO DAS DECISÕES", "RESPONSABILIDADES SINTÉTICAS", "CONDIÇÕES DE RECEBIMENTO",
        "INSPEÇÃO DOS MATERIAIS", "CONTROLE DO CRONOGRAMA", "REGISTROS DE ENTREGA", "MEDIÇÃO DO ENSAIO",
        "CONFERÊNCIA DOS TOTAIS", "TRATAMENTO DAS DIVERGÊNCIAS", "REGISTRO DE AJUSTES",
        "PRESERVAÇÃO DAS VERSÕES", "REVISÃO DOS ANEXOS", "ORDEM DOS DOCUMENTOS",
        "REFERÊNCIAS SINTÉTICAS", "NUMERAÇÃO DAS CLÁUSULAS", "ENCERRAMENTO DO ENSAIO", "DISPOSIÇÕES DE QA"), 1)]
    edital_pieces = [SYNTHETIC_NOTICE, "## Índice sintético\n",
                      "\n".join(f"Índice de teste para {title}" for title in edital_titles)]
    edital_markers, edital_links = [], []
    for i, title in enumerate(edital_titles, 1):
        marker = f"C05-CLAUSE-{i:02d}"
        edital_markers.append(marker)
        edital_pieces.append(
            f"\n## {title}\n\n{marker} Esta cláusula pertence ao edital fictício usado no teste "
            "de preservação estrutural. A equipe registra a origem de cada informação, confere "
            "o vínculo com os anexos e mantém o histórico de decisões do cenário. O texto "
            "descreve operações de ensaio, sem estabelecer obrigações para pessoas ou órgãos "
            "reais. As remissões ao final apontam para os próprios títulos deste documento. "
            "Uma alteração de página não pode excluir a cláusula, trocar sua ordem ou deslocar "
            "seus registros para outro documento. A análise posterior verificará cada marcador, "
            "o destino das remissões, a integridade dos links e a continuidade da numeração.\n")
        if i % 4 == 0:
            link = f"https://example.invalid/edital/anexo-{i:02d}"
            edital_links.append(link)
            edital_pieces.append(f"[Anexo sintético {i:02d}]({link})\n")
    edital = "\n".join(edital_pieces)
    arp = SYNTHETIC_NOTICE + "## Ata de registro de preços de ensaio\n\nC06-ARP-ONLY Instrumento condicional sintético.\n\n" + dfd_table
    positive, negative = {"modelo_execucao": "Sistema de Registro de Preços"}, {"modelo_execucao": "Fornecimento sem SRP"}
    bundle = {"edital": SYNTHETIC_NOTICE + "C06-EDITAL-ONLY Edital do cenário condicional.", "arp": arp}
    if ("arp" not in config.exportaveis_do_processo(positive, bundle)
            or "arp" in config.exportaveis_do_processo(negative, bundle)):
        raise ValueError("Canonical conditional ARP rule differs from the scenario")

    logo = _logo_sintetico()
    live_secretary = copy.deepcopy(institution["secretary"])
    live_secretary["cabecalho_img"] = base64.b64encode(logo).decode("ascii")
    selected, selected_origin = contexto.resolver_identidade(
        [institution["municipal"], live_secretary], live_secretary["id"])
    frozen_brand = copy.deepcopy(selected)
    fallback, fallback_origin = contexto.resolver_identidade(
        [institution["municipal"], {"id": "qa-no-own-identity"}], "qa-no-own-identity")
    table_210, oracle_210 = tabela_sintetica()
    checked_table, checked_oracle = tabela_com_oraculo(itens_sinteticos())
    if checked_table != table_210:
        raise ValueError("Full-corpus calculation changed the existing 210-item fixture source")
    oracle_210.update(checked_oracle)
    objective_target = "## 1. Prazo de teste\n\nC09-ALVO O prazo de teste está indeterminado.\n\n## 2. Assinaturas preservadas\n"
    findings = [
        {"id": "C09-F01", "nature": "objetiva", "recommendation": "Definir 30 dias no trecho C09-ALVO.",
         "expected_disposition": "PROPOSAL_REQUIRED_NOT_APPLIED", "target_marker": "C09-ALVO"},
        {"id": "C09-F02", "nature": "ambigua", "recommendation": "Melhorar a adequação do documento.",
         "expected_disposition": "PENDING_CLARIFICATION"},
        {"id": "C09-F03", "nature": "conflitante", "recommendation": "Remover assinaturas institucionais preservadas.",
         "expected_disposition": "PENDING_CONFLICT_REVIEW"}]
    opinion = SYNTHETIC_NOTICE + "## Parecer de teste sem efeito jurídico\n\n" + "\n\n".join(
        f"{finding['id']} {finding['nature']}: {finding['recommendation']}" for finding in findings)
    external = (SYNTHETIC_NOTICE + "## Modelo de terceiro simulado\n\nC10-EXTERNAL "
        "Modelo original de QA com tabela aninhada, célula mesclada e campo de página.\n\n"
        "[Referência declarada](https://example.invalid/modelo-externo)\n\n"
        "| Quadro externo | Recursos declarados |\n|---|---|\n| C10-OUTER | Conteúdo da célula externa |\n")
    legacy = SYNTHETIC_NOTICE + "# Original legado simulado\n\nC11-LEGACY-ORIGINAL Texto protegido de teste.\n"
    revised = legacy.replace("C11-LEGACY-ORIGINAL Texto protegido", "C11-REVISION-NEW Nova revisão")
    tenant_b = instituicao_sintetica("B")
    return [
        case("C01", "dfd-5-itens", [variant("dfd", "DFD sintético com cinco itens", dfd,
             oracle=dfd_oracle, brand=branding, sign=signature)]),
        case("C02", "etp-longo", [variant("etp", "ETP sintético para paginação longa", etp,
             oracle=etp_oracle, brand=branding, sign=signature)]),
        case("C03", "mapa-riscos-mesclas", [variant("mapa-riscos", "Mapa de Riscos", risks,
             qa="risk_merges", oracle={"markers": risk_codes, "horizontal_span": 6,
                 "vertical_merge_restarts": 3, "vertical_merge_continuations": 9,
                 "explicit_page_break_before": 1}, brand=branding)]),
        case("C04", "tr-30-itens", [variant("tr", "TR sintético com trinta itens", tr,
             oracle=tr_oracle, brand=branding, sign=signature)]),
        case("C05", "edital-remissoes", [variant("edital", "Edital sintético com remissões", edital,
             qa="bookmarks", oracle={"markers": edital_markers, "clause_titles": edital_titles,
                 "external_links": edital_links, "bookmarks": [f"gc_c05_{i:02d}" for i in range(1, 33)]},
             brand=branding)]),
        case("C06", "arp-condicional", [
             variant("arp-srp", "ARP sintética aplicável ao cenário SRP", arp,
                 oracle={**dfd_money_oracle, "markers": ["C06-ARP-ONLY"], "conditional_arp_expected": True}),
             {**variant("dossie-sem-srp", "Dossiê sintético sem SRP", bundle["edital"],
                 oracle={"markers": ["C06-EDITAL-ONLY"], "forbidden_text": ["C06-ARP-ONLY"],
                         "conditional_arp_expected": False}), "consolidated": {"documents": bundle, "dados": negative}}],
             workflow=["OPERATIONAL_CONDITIONAL_EMISSION_NOT_TESTED"]),
        case("C07", "timbrado-snapshot", [
             variant("secretaria", "Timbrado e snapshot sintéticos da secretaria",
                 SYNTHETIC_NOTICE + "**SECRETARIA SINTÉTICA A**\n\n"
                 "C07-SECRETARIA Identidade própria capturada antes de uma alteração de cadastro.",
                 brand=frozen_brand, sign=signature, oracle={"markers": ["C07-SECRETARIA"],
                     "required_text": ["SECRETARIA SINTÉTICA A"],
                     "institutional_origin": selected_origin, "branding_snapshot": frozen_brand,
                     "logo_sha256": hashlib.sha256(logo).hexdigest(),
                     "signatarios_snapshot": institution["signatarios_snapshot"],
                     "signature_names": institution["signature_names"]}),
             variant("municipio", "Timbrado sintético herdado do município",
                 SYNTHETIC_NOTICE + "C07-MUNICIPIO Herança municipal sem identidade própria.",
                 brand=fallback, oracle={"markers": ["C07-MUNICIPIO"], "institutional_origin": fallback_origin,
                                        "branding_snapshot": copy.deepcopy(fallback)})],
             workflow=["APPROVED_SNAPSHOT_AFTER_REGISTRY_CHANGE_NOT_TESTED"]),
        case("C08", "tabela-210-itens", [variant("planilha", "Planilha sintética com duzentos e dez itens",
             SYNTHETIC_NOTICE + table_210, oracle=oracle_210)]),
        case("C09", "parecer-achados", [variant("parecer", "Parecer sintético para correção endereçada", opinion,
             oracle={"markers": [finding["id"] for finding in findings], "findings": findings,
                     "source_markdown_hash": blocos.hash_texto(objective_target)})],
             workflow=["CANONICAL_PATCH_DIFF_REVALIDATION_NOT_TESTED"], support={"alvo-sintetico.md": objective_target}),
        case("C10", "modelo-externo-simulado", [variant("modelo", "Modelo externo exclusivamente sintético", external,
             qa="external_model", brand=branding,
             oracle={"markers": ["C10-EXTERNAL", "C10-OUTER", "C10-NEST-A", "C10-NEST-B", "C10-PAGE-2"],
                     "external_links": ["https://example.invalid/modelo-externo"],
                     "declared_features": ["nested_table", "horizontal_merge", "header", "footer", "page_field", "page_break"]})]),
        case("C11", "legado-historico", [
             variant("original", "Original histórico simulado", legacy,
                 oracle={"markers": ["C11-LEGACY-ORIGINAL"], "forbidden_text": ["C11-REVISION-NEW"]}),
             variant("nova-revisao", "Nova revisão sintética separada", revised,
                 oracle={"markers": ["C11-REVISION-NEW"], "forbidden_text": ["C11-LEGACY-ORIGINAL"]})],
             workflow=["APPROVAL_AND_IMMUTABILITY_NOT_TESTED"], support={"original-legado.md": legacy}),
        case("C12", "dois-tenants-secretarias", [
             variant("tenant-a", "Documento sintético do tenant A", SYNTHETIC_NOTICE + "QA-TENANT-A-ONLY",
                 brand=branding, oracle={"markers": ["QA-TENANT-A-ONLY"], "forbidden_text": ["QA-TENANT-B-ONLY"],
                     "tenant_id": institution["tenant_id"], "secretaria_id": institution["secretaria_id"]}),
             variant("tenant-b", "Documento sintético do tenant B", SYNTHETIC_NOTICE + "QA-TENANT-B-ONLY",
                 brand=tenant_b["branding_snapshot"],
                 oracle={"markers": ["QA-TENANT-B-ONLY"], "forbidden_text": ["QA-TENANT-A-ONLY"],
                     "tenant_id": tenant_b["tenant_id"], "secretaria_id": tenant_b["secretaria_id"]})],
             workflow=["JWT_RLS_STORAGE_CACHE_BRIDGE_ISOLATION_NOT_TESTED"]),
    ]


def _verificar_entrada_docx(data: bytes, oracle: dict) -> dict:
    """Check authored inputs against their source oracles, not WordCraft fidelity."""
    inspection = inspect_full_docx(data)
    text = inspection["text"]
    if inspection["page_twips"] != {"w": "11906", "h": "16838"}:
        raise ValueError("Full corpus requires canonical A4 inputs")
    for marker in oracle.get("markers", []):
        if text.count(marker) != 1:
            raise ValueError("Generated input lost or duplicated an independent marker")
    if any(value in text for value in oracle.get("forbidden_text", [])):
        raise ValueError("Generated input contains a forbidden scenario marker")
    for name in oracle.get("signature_names", []):
        if name not in text:
            raise ValueError("Generated input lost a synthetic snapshot signatory")
    if any(value not in text for value in oracle.get("required_text", [])):
        raise ValueError("Generated input lost required institutional text")
    for code in oracle.get("codes", []):
        if text.count(code) != 1:
            raise ValueError("Generated input lost or duplicated an item code")
    if "expected_money" in oracle:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
        rows = [[("".join(node.text or "" for node in cell.iter(W + "t"))).strip()
                 for cell in row.findall(W + "tc")] for row in root.iter(W + "tr")]
        expected_codes = set(oracle["codes"])
        if [row[0] for row in rows if row and row[0] in expected_codes] != oracle["codes"]:
            raise ValueError("DOCX item order differs from the independent source")
        for source, total in zip(oracle["source_items"], oracle["expected_money"]):
            matches = [row for row in rows if row and row[0] == source["codigo"]]
            unit = _moeda_oraculo(Decimal(source["valor_unitario"].replace(",", ".")))
            if len(matches) != 1 or len(matches[0]) < 6 or matches[0][4:6] != [unit, total]:
                raise ValueError("DOCX monetary cells differ from the independent Decimal oracle")
            if (matches[0][1] != source["descricao"] or matches[0][2] != source["unidade"]
                    or Decimal(matches[0][3].replace(",", ".")) != Decimal(source["quantidade"])):
                raise ValueError("DOCX item identity or quantity differs from its source")
        totals = [row for row in rows if "VALOR GLOBAL" in row]
        if len(totals) != 1 or len(totals[0]) < 6 or totals[0][5] != oracle["global_total_display"]:
            raise ValueError("DOCX global total differs from the independent Decimal oracle")
    expected_links = oracle.get("external_links")
    if expected_links is not None and sorted(text_blocks(data)["links"]) != sorted(expected_links):
        raise ValueError("DOCX hyperlinks differ from the declared synthetic references")
    if "bookmarks" in oracle:
        if inspection["bookmarks"] != oracle["bookmarks"] or inspection["internal_links"] != oracle["bookmarks"]:
            raise ValueError("DOCX internal references differ from the expected clause anchors")
    if "logo_sha256" in oracle and oracle["logo_sha256"] not in inspection["media_sha256"]:
        raise ValueError("DOCX lost the captured synthetic branding asset")
    if "horizontal_span" in oracle and oracle["horizontal_span"] not in inspection["horizontal_spans"]:
        raise ValueError("DOCX lost the declared horizontal merge")
    for key in ("vertical_merge_restarts", "vertical_merge_continuations"):
        if key in oracle and inspection[key] != oracle[key]:
            raise ValueError("DOCX vertical merges differ from the declared fixture")
    inspection.pop("text")
    return inspection


def generate_full(output: Path, *, pdf=False) -> dict:
    """Generate all twelve scenario seeds; no claim of twelve passing gates."""
    from src.export import _docx_em_pdf, gerar_docx_consolidado
    output = output.resolve()
    protected = (ROOT / "tests/fixtures/wordcraft").resolve()
    if output.is_relative_to(protected) or output == ROOT or (output.exists() and any(output.iterdir())):
        raise ValueError("Use a fresh full-corpus directory outside the immutable G1 inputs")
    if pdf and motor_pdf() != "libreoffice":
        raise ValueError("Full PDF corpus requires the real institutional LibreOffice engine")
    manifest = {"schema": FULL_SCHEMA, "synthetic_only": True,
                "status": "INPUTS_GENERATED_NOT_ROUNDTRIP_APPROVED",
                "scope": "QA_INPUTS_ONLY_NOT_APPROVAL_OR_RLS_EVIDENCE", "cases": [], "files": []}
    writes = []
    for case in full_corpus_cases():
        entry = {"id": case["id"], "name": case["name"], "variants": [],
                 "workflow_acceptance": {name: "NOT_TESTED" for name in case["workflow_acceptance"]}}
        for fixture in case["variants"]:
            prefix = case["id"] + "/" + fixture["name"]
            if "consolidated" in fixture:
                source = fixture["consolidated"]
                data = gerar_docx_consolidado(source["documents"], dados=source["dados"])
            else:
                data = gerar_docx(fixture["title"], fixture["markdown"],
                                  fixture["branding"], fixture["signatures"])
            if fixture["qa"]:
                data = _enriquecer_qa(data, fixture["qa"],
                                      clause_titles=fixture["oracle"].get("clause_titles", []))
            inspection = _verificar_entrada_docx(data, fixture["oracle"])
            oracle = {"synthetic_only": True, "status": "EXPECTED_ONLY_NOT_PRODUCT_ACCEPTANCE",
                      **fixture["oracle"]}
            writes.extend([(prefix + ".docx", data),
                           (prefix + ".md", (fixture["markdown"] + fixture["signatures"]).encode()),
                           (prefix + ".oracle.json", (json.dumps(oracle, ensure_ascii=False, indent=2) + "\n").encode())])
            variant = {"name": fixture["name"], "docx": prefix + ".docx",
                       "oracle": prefix + ".oracle.json", "sha256_docx": hashlib.sha256(data).hexdigest(),
                       "origin": "canonical_export_plus_python_docx_QA" if fixture["qa"] else "canonical_product_export",
                       "qa_enrichment": fixture["qa"], "input_inspection": inspection,
                       "fidelity": {"wordcraft_roundtrip": "NOT_RUN", "visual_review": "PENDING",
                                    "libreoffice_pdf": "NOT_REQUESTED"}}
            if pdf:
                rendered = _docx_em_pdf(data)
                if not rendered or not rendered.startswith(b"%PDF-") or motor_pdf() != "libreoffice":
                    raise ValueError("LibreOffice failed; no PDF fallback can approve this corpus")
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(rendered))
                extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
                if any(marker not in extracted for marker in fixture["oracle"].get("codes", [])):
                    raise ValueError("LibreOffice PDF lost a synthetic item code")
                writes.append((prefix + ".pdf", rendered))
                variant.update(pdf=prefix + ".pdf", sha256_pdf=hashlib.sha256(rendered).hexdigest(),
                               pages=len(reader.pages), pdf_engine="libreoffice")
                variant["fidelity"]["libreoffice_pdf"] = "GENERATED_VISUAL_REVIEW_PENDING"
            entry["variants"].append(variant)
        for name, content in case["support"].items():
            writes.append((case["id"] + "/" + name, content.encode()))
        manifest["cases"].append(entry)
    if {entry["id"] for entry in manifest["cases"]} != {f"C{i:02d}" for i in range(1, 13)}:
        raise ValueError("Full corpus must declare all twelve scenario seeds")
    # Complete generation/validation before creating the target directory.
    output.mkdir(parents=True, exist_ok=True)
    for relative, data in writes:
        target = output / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(data)
        manifest["files"].append({"path": relative, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()})
    (output / "full-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def verify_full_inputs(inputs: Path) -> dict:
    """Verify every full-corpus byte; never mark roundtrip/governance as passed."""
    inputs = inputs.resolve(strict=True)
    manifest = json.loads((inputs / "full-manifest.json").read_text(encoding="utf-8"))
    if (manifest.get("schema") != FULL_SCHEMA or manifest.get("synthetic_only") is not True
            or manifest.get("status") != "INPUTS_GENERATED_NOT_ROUNDTRIP_APPROVED"
            or len(manifest.get("cases", [])) != 12
            or {case["id"] for case in manifest["cases"]} != {f"C{i:02d}" for i in range(1, 13)}):
        raise ValueError("Invalid full-corpus manifest")
    seen = set()
    for item in manifest["files"]:
        name = item["path"]
        if (not isinstance(name, str) or "\\" in name or ":" in name or name.startswith("/")
                or any(part in {"", ".", ".."} for part in name.split("/")) or name in seen):
            raise ValueError("Invalid full-corpus file path")
        seen.add(name)
        target = (inputs / name).resolve(strict=True)
        if not target.is_relative_to(inputs) or not 0 < target.stat().st_size <= 8 * 1024 * 1024:
            raise ValueError("Full-corpus input escaped its bounded namespace")
        data = target.read_bytes()
        if len(data) != item["bytes"] or hashlib.sha256(data).hexdigest() != item["sha256"]:
            raise ValueError("Full-corpus fixture digest differs")
    for case in manifest["cases"]:
        for variant in case["variants"]:
            if variant["docx"] not in seen or variant["oracle"] not in seen:
                raise ValueError("Full-corpus case lacks an inventoried input")
            oracle = json.loads((inputs / variant["oracle"]).read_text(encoding="utf-8"))
            _verificar_entrada_docx((inputs / variant["docx"]).read_bytes(), oracle)
    actual = {path.relative_to(inputs).as_posix() for path in inputs.rglob("*")
              if path.is_file() and path.name != "full-manifest.json"}
    if actual != seen:
        raise ValueError("Uninventoried file in full-corpus directory")
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--pdf", action="store_true", help="Require the institutional LibreOffice engine")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--compare-roundtrip", action="store_true")
    mode.add_argument("--full-corpus", action="store_true", help="Generate separate G3 scenario seeds C01-C12")
    parser.add_argument("--inputs", type=Path)
    args = parser.parse_args()
    if args.compare_roundtrip:
        if args.inputs is None:
            parser.error("--compare-roundtrip needs --inputs")
        result = compare_roundtrip(args.inputs, args.out)
    else:
        result = generate_full(args.out, pdf=args.pdf) if args.full_corpus else generate(args.out, pdf=args.pdf)
    print(json.dumps(result, ensure_ascii=False, indent=2))
